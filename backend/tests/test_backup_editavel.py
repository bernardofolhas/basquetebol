import json
from io import BytesIO
from uuid import uuid4
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from auth import encriptar_password
from database import SessionLocal
from models.clube import Clube
from models.equipa import Equipa, EscalaoEnum, GeneroEnum
from models.treinador import Treinador, treinador_equipa
from models.utilizador import PerfilEnum, Utilizador
from routers.exportacoes import _criar_excel


@pytest.fixture
def cliente_treinador_backup(client):
    marcador = uuid4().hex[:8]
    db = SessionLocal()
    clube = db.query(Clube).order_by(Clube.id).first()
    clube_id = clube.id
    equipa = Equipa(
        clube_id=clube_id,
        nome=f"Equipa Backup {marcador}",
        escalao=EscalaoEnum.sub18,
        genero=GeneroEnum.misto,
    )
    treinador = Treinador(clube_id=clube_id, nome=f"Treinador Backup {marcador}")
    db.add_all([equipa, treinador])
    db.flush()
    treinador.equipas.append(equipa)
    utilizador = Utilizador(
        clube_id=clube_id,
        treinador_id=treinador.id,
        nome=treinador.nome,
        username=f"backup_{marcador}",
        email=f"backup_{marcador}@example.test",
        password_hash=encriptar_password("Backup1234!"),
        perfil=PerfilEnum.treinador,
        ativo=True,
    )
    db.add(utilizador)
    db.commit()
    equipa_id = equipa.id
    treinador_id = treinador.id
    utilizador_id = utilizador.id
    username = utilizador.username
    db.close()

    resposta = client.post(
        "/login-form",
        data={"username": username, "password": "Backup1234!"},
        follow_redirects=False,
    )
    assert resposta.status_code in {302, 303}
    equipa_sem_acesso = Equipa(
        clube_id=clube_id,
        nome=f"Equipa Sem Acesso {marcador}",
        escalao=EscalaoEnum.sub16,
        genero=GeneroEnum.misto,
    )
    db = SessionLocal()
    db.add(equipa_sem_acesso)
    db.commit()
    equipa_sem_acesso_id = equipa_sem_acesso.id
    db.close()
    yield client, equipa_id, equipa_sem_acesso_id

    db = SessionLocal()
    try:
        db.query(Utilizador).filter(Utilizador.id == utilizador_id).delete(synchronize_session=False)
        db.execute(treinador_equipa.delete().where(treinador_equipa.c.treinador_id == treinador_id))
        db.query(Treinador).filter(Treinador.id == treinador_id).delete(synchronize_session=False)
        db.query(Equipa).filter(Equipa.id == equipa_id).delete(synchronize_session=False)
        db.query(Equipa).filter(Equipa.id == equipa_sem_acesso_id).delete(synchronize_session=False)
        db.commit()
    finally:
        db.close()


@pytest.fixture
def cliente_coordenador_backup(client):
    marcador = uuid4().hex[:8]
    db = SessionLocal()
    clube = db.query(Clube).order_by(Clube.id).first()
    utilizador = Utilizador(
        clube_id=clube.id,
        nome=f"Coordenador Backup {marcador}",
        username=f"coord_backup_{marcador}",
        email=f"coord_backup_{marcador}@example.test",
        password_hash=encriptar_password("Backup1234!"),
        perfil=PerfilEnum.coordenador,
        ativo=True,
    )
    db.add(utilizador)
    db.commit()
    utilizador_id = utilizador.id
    username = utilizador.username
    db.close()

    resposta = client.post(
        "/login-form",
        data={"username": username, "password": "Backup1234!"},
        follow_redirects=False,
    )
    assert resposta.status_code in {302, 303}
    yield client

    db = SessionLocal()
    try:
        db.query(Utilizador).filter(Utilizador.id == utilizador_id).delete(synchronize_session=False)
        db.commit()
    finally:
        db.close()


def test_excel_e_editavel_seguro_e_preserva_integral_no_json():
    utilizador = Utilizador(
        nome="Teste",
        username="=FORMULA",
        email="teste@example.test",
        password_hash="segredo",
        perfil="treinador",
    )
    conteudo = _criar_excel("Utilizadores", [("utilizadores", [utilizador])])
    livro = load_workbook(BytesIO(conteudo))
    folha = livro["utilizadores"]
    cabecalhos = [celula.value for celula in folha[1]]

    assert "password_hash" not in cabecalhos
    assert "sessao_versao" not in cabecalhos
    assert folha.cell(2, cabecalhos.index("username") + 1).value == "'=FORMULA"
    assert folha.freeze_panes == "A2"


def test_admin_descarrega_backup_do_clube_organizado_por_equipa(client_admin):
    resposta = client_admin.get("/exportacoes/clube/dados")
    assert resposta.status_code == 200
    assert resposta.headers["content-type"].startswith("application/zip")

    with ZipFile(BytesIO(resposta.content)) as zipado:
        nomes = zipado.namelist()
        manifesto = json.loads(zipado.read("MANIFESTO.json"))
        assert manifesto["escopo"] == "clube"
        assert "00_Clube/identificacao_do_clube.xlsx" in nomes
        assert "02_Planeamento_do_clube/planeamento_do_clube.xlsx" in nomes
        assert "03_Espacos_e_horarios/espacos_horarios_e_atribuicoes.xlsx" in nomes
        assert all(any(nome.startswith(equipa["pasta"] + "/") for nome in nomes) for equipa in manifesto["equipas"])

        for nome in (item for item in nomes if item.endswith(".xlsx")):
            livro = load_workbook(BytesIO(zipado.read(nome)), read_only=True)
            assert livro.sheetnames
        for nome in (item for item in nomes if item.endswith("_integral.json")):
            assert isinstance(json.loads(zipado.read(nome)), dict)

        utilizadores_json = json.loads(
            zipado.read("01_Utilizadores/utilizadores_e_equipa_tecnica_integral.json")
        )
        assert all("password_hash" not in linha for linha in utilizadores_json["utilizadores"])
        assert all("sessao_versao" not in linha for linha in utilizadores_json["utilizadores"])


def test_coordenador_tambem_descarrega_backup_do_clube(cliente_coordenador_backup):
    resposta = cliente_coordenador_backup.get("/exportacoes/clube/dados")
    assert resposta.status_code == 200
    with ZipFile(BytesIO(resposta.content)) as zipado:
        manifesto = json.loads(zipado.read("MANIFESTO.json"))
        assert manifesto["escopo"] == "clube"


def test_treinador_descarrega_backup_apenas_da_sua_equipa(cliente_treinador_backup):
    client, equipa_id, equipa_sem_acesso_id = cliente_treinador_backup
    resposta = client.get(f"/exportacoes/equipas/{equipa_id}/dados")
    assert resposta.status_code == 200
    with ZipFile(BytesIO(resposta.content)) as zipado:
        manifesto = json.loads(zipado.read("MANIFESTO.json"))
        assert manifesto["escopo"] == "equipa"
        assert [item["id"] for item in manifesto["equipas"]] == [equipa_id]
        assert not any(nome.startswith("06_Equipas/") for nome in zipado.namelist())
        assert "00_Clube/identificacao_do_clube.xlsx" in zipado.namelist()

    negado = client.get(f"/exportacoes/equipas/{equipa_sem_acesso_id}/dados", follow_redirects=False)
    assert negado.status_code in {302, 303, 403}


def test_treinador_nao_pode_descarregar_backup_total_do_clube(cliente_treinador_backup):
    client, _, _ = cliente_treinador_backup
    resposta = client.get("/exportacoes/clube/dados", follow_redirects=False)
    assert resposta.status_code in {302, 303, 403}
