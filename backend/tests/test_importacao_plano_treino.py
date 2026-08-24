from io import BytesIO
from datetime import date
from pathlib import Path
from zipfile import ZipFile
from types import SimpleNamespace

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from weasyprint import HTML

from services.importacao_plano_treino_service import (
    ErroImportacaoPlano,
    extrair_texto,
    interpretar_texto,
    normalizar_categoria,
    sincronizar_minutos_categorias,
    validar_ficheiro,
)
from template_config import Jinja2Templates
from paths import TEMPLATES_DIR
from database import SessionLocal
from models.treino import ImportacaoPlanoTreino, PlanoTreino
from models.utilizador import Utilizador
from models.equipa import Equipa, EscalaoEnum, GeneroEnum
from models.planeamento import Macrociclo, Mesociclo, Microciclo, SessaoTreino


TEXTO_PLANO = """Plano de treino
Data: 24/08/2026
Hora: 18:30 - 20:00
Local: Pavilhão Fidelidade
Escalão: Sub-23
Objetivos: Transição ofensiva e recuperação defensiva
Bloco A
18:30 | 15 min | Finalização após contra-ataque | TIO | RPE 6
18:45 | Defesa do bloqueio direto | TID | 20 minutos | RPE 7
Bloco B
19:05 | 25 min | Jogo condicionado | TCO | RPE 8
"""


def test_interpretacao_reconhece_cabecalho_exercicios_e_minutos_por_categoria():
    dados = interpretar_texto(TEXTO_PLANO)

    assert dados["cabecalho"] == {
        "data": "2026-08-24",
        "hora_inicio": "18:30",
        "hora_fim": "20:00",
        "local": "Pavilhão Fidelidade",
        "escalao": "Sub-23",
        "objectivos": "Transição ofensiva e recuperação defensiva",
        "rpe_geral": "",
        "volume_minutos": "60",
    }
    assert [item["nome"] for item in dados["exercicios"]] == [
        "Finalização após contra-ataque",
        "Defesa do bloqueio direto",
        "Jogo condicionado",
    ]
    assert dados["totais_categoria"]["TIO"] == 15
    assert dados["totais_categoria"]["TID"] == 20
    assert dados["totais_categoria"]["TCO"] == 25


@pytest.mark.parametrize(
    ("valor", "esperado"),
    [
        ("TIO", "TIO"),
        ("técnica individual defensiva", "TID"),
        ("Tática individual ofensiva", "TaIO"),
        ("preparação física", "FIS"),
        ("psicológico", "PSI"),
    ],
)
def test_normalizacao_categorias_aceita_siglas_e_descricoes(valor, esperado):
    assert normalizar_categoria(valor) == esperado


def test_extracao_de_word_inclui_paragrafos_e_tabelas():
    documento = Document()
    documento.add_paragraph("Data: 24/08/2026")
    tabela = documento.add_table(rows=1, cols=3)
    tabela.rows[0].cells[0].text = "Passe e corte"
    tabela.rows[0].cells[1].text = "TIO"
    tabela.rows[0].cells[2].text = "15 min"
    memoria = BytesIO()
    documento.save(memoria)

    texto, confianca, avisos = extrair_texto("plano.docx", memoria.getvalue())

    assert "Data: 24/08/2026" in texto
    assert "Passe e corte | TIO | 15 min" in texto
    assert confianca == 100
    assert avisos == []


def test_extracao_de_excel_preserva_linhas_editaveis():
    livro = Workbook()
    folha = livro.active
    folha.title = "Treino 12"
    folha.append(["Exercício", "Categoria", "Minutos"])
    folha.append(["1x1", "TID", 12])
    memoria = BytesIO()
    livro.save(memoria)

    texto, confianca, avisos = extrair_texto("plano.xlsx", memoria.getvalue())

    assert "Folha: Treino 12" in texto
    assert "1x1 | TID | 12" in texto
    assert confianca == 100
    assert avisos == []


def test_extracao_de_pdf_digitado_preserva_texto():
    pdf = HTML(string="<p>Plano de treino</p><p>Finalização | TIO | 15 min</p>").write_pdf()

    texto, confianca, avisos = extrair_texto("plano.pdf", pdf)

    assert "Plano de treino" in texto
    assert "Finalização" in texto
    assert confianca == 100
    assert avisos == []


def test_fotografia_usa_ocr_e_devolve_confianca(monkeypatch):
    import pytesseract

    monkeypatch.setattr(
        pytesseract,
        "image_to_data",
        lambda *args, **kwargs: {
            "text": ["Passe", "TIO", "15", "min"],
            "conf": ["91", "94", "92", "90"],
            "block_num": [1, 1, 1, 1],
            "par_num": [1, 1, 1, 1],
            "line_num": [1, 1, 1, 1],
        },
    )
    memoria = BytesIO()
    Image.new("RGB", (1000, 700), "white").save(memoria, format="PNG")

    texto, confianca, avisos = extrair_texto("fotografia.png", memoria.getvalue())

    assert texto == "Passe TIO 15 min"
    assert confianca == pytest.approx(91.75)
    assert avisos == []


def test_interpretacao_de_tabela_preserva_organizacao_e_coaching_points():
    dados = interpretar_texto(
        "Hora | Minutos | Exercício | Categoria | RPE | Organização | Coaching Points\n"
        "18:30 | 12 min | Passe e corte | TIO | 6 | 3 filas na linha final | Cortar de frente para a bola"
    )

    assert dados["exercicios"] == [{
        "bloco": "A",
        "nome": "Passe e corte",
        "categoria": "TIO",
        "hora": "18:30",
        "minutos": 12,
        "rpe": "6",
        "organizacao": "3 filas na linha final",
        "coaching_points": "Cortar de frente para a bola",
        "confianca": "alta",
    }]


def test_validacao_rejeita_extensao_conteudo_e_tamanho_invalidos():
    with pytest.raises(ErroImportacaoPlano):
        validar_ficheiro("plano.txt", "text/plain", b"texto")
    with pytest.raises(ErroImportacaoPlano):
        validar_ficheiro("plano.pdf", "application/pdf", b"nao e pdf")
    with pytest.raises(ErroImportacaoPlano):
        validar_ficheiro("plano.png", "image/png", b"")
    with pytest.raises(ErroImportacaoPlano):
        validar_ficheiro("plano.docx", "application/octet-stream", b"PKdocumento-danificado")


def test_data_impossivel_nao_e_assumida():
    dados = interpretar_texto("Data: 31/02/2026\nPasse | TIO | 15 min")

    assert dados["cabecalho"]["data"] == ""
    assert any("data" in aviso.lower() for aviso in dados["avisos"])


def test_sincronizacao_distribui_minutos_da_sessao():
    sessao = SimpleNamespace()
    plano = SimpleNamespace(
        sessao=sessao,
        exercicios=[
            SimpleNamespace(categoria="TIO", minutos=15),
            SimpleNamespace(categoria="técnica individual ofensiva", minutos=5),
            SimpleNamespace(categoria="TID", minutos=20),
            SimpleNamespace(categoria="FIS", minutos=None),
        ],
    )

    totais = sincronizar_minutos_categorias(plano)

    assert totais["TIO"] == 20
    assert totais["TID"] == 20
    assert sessao.min_tecnica_individual_ofensiva == 20
    assert sessao.min_tecnica_individual_defensiva == 20
    assert sessao.min_fisicos == 0


def test_templates_de_upload_e_revisao_sao_validos():
    templates = Jinja2Templates(directory=str(TEMPLATES_DIR))
    templates.env.get_template("treino/plano.html")
    revisao = templates.env.get_template("treino/importacao_rever.html")
    conteudo_plano = (Path(TEMPLATES_DIR) / "treino" / "plano.html").read_text(encoding="utf-8")

    assert "multipart/form-data" in conteudo_plano
    assert revisao is not None


def test_upload_autenticado_cria_rascunho_e_abre_revisao(client_admin):
    db = SessionLocal()
    importacao_id = None
    equipa_id = None
    macro_id = None
    try:
        admin = db.query(Utilizador).execution_options(ignorar_clube=True).filter(Utilizador.username == "admin").first()
        if not admin or not admin.clube_id:
            pytest.skip("A instalação de teste não tem o administrador de clube esperado.")
        equipa = Equipa(
            clube_id=admin.clube_id,
            nome="Equipa temporária importação",
            escalao=EscalaoEnum.sub18,
            genero=GeneroEnum.misto,
        )
        db.add(equipa)
        db.flush()
        equipa_id = equipa.id
        macro = Macrociclo(clube_id=admin.clube_id, equipa_id=equipa.id, numero=1, data_inicio=date(2026, 8, 1), data_fim=date(2026, 8, 31))
        db.add(macro)
        db.flush()
        macro_id = macro.id
        meso = Mesociclo(clube_id=admin.clube_id, macrociclo_id=macro.id, numero=1, data_inicio=date(2026, 8, 1), data_fim=date(2026, 8, 31))
        db.add(meso)
        db.flush()
        micro = Microciclo(clube_id=admin.clube_id, mesociclo_id=meso.id, numero=1, data_inicio=date(2026, 8, 24), data_fim=date(2026, 8, 30))
        db.add(micro)
        db.flush()
        sessao = SessaoTreino(clube_id=admin.clube_id, microciclo_id=micro.id, data=date(2026, 8, 24), e_jogo=False)
        db.add(sessao)
        db.flush()
        plano = PlanoTreino(clube_id=admin.clube_id, sessao_id=sessao.id)
        db.add(plano)
        db.commit()

        pagina_plano = client_admin.get(f"/treino/sessoes/{sessao.id}")
        assert pagina_plano.status_code == 200
        assert "Importar plano" in pagina_plano.text
        assert "multipart/form-data" in pagina_plano.text

        documento = Document()
        documento.add_paragraph(TEXTO_PLANO)
        memoria = BytesIO()
        documento.save(memoria)
        resposta = client_admin.post(
            f"/treino/planos/{plano.id}/importacoes",
            files={"ficheiro": ("plano-teste.docx", memoria.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            follow_redirects=False,
        )

        assert resposta.status_code == 303
        assert resposta.headers["location"].startswith("/treino/importacoes/")
        importacao_id = int(resposta.headers["location"].rsplit("/", 1)[-1])
        revisao = client_admin.get(resposta.headers["location"])
        assert revisao.status_code == 200
        assert "Rever plano importado" in revisao.text
        assert "Finalização após contra-ataque" in revisao.text
        assert "Aplicar ao plano" in revisao.text

        original = client_admin.get(f"/treino/importacoes/{importacao_id}/ficheiro")
        assert original.status_code == 200
        assert original.content == memoria.getvalue()
        assert original.headers["cache-control"] == "private, no-store"
        assert original.headers["x-content-type-options"] == "nosniff"
        aplicacao = client_admin.post(
            f"/treino/importacoes/{importacao_id}/aplicar",
            data={
                "data": "2026-08-24",
                "hora_inicio": "18:30",
                "hora_fim": "20:00",
                "local": "Pavilhão de teste",
                "escalao": "Sub-18",
                "rpe_geral": "7",
                "volume_minutos": "35",
                "objectivos": "Validar a importação",
                "modo": "adicionar",
                "exercicio_bloco": ["A", "A"],
                "exercicio_nome": ["Finalização", "Defesa 1x1"],
                "exercicio_categoria": ["TIO", "TID"],
                "exercicio_minutos": ["15", "20"],
                "exercicio_hora": ["18:30", "18:45"],
                "exercicio_rpe": ["6", "7"],
                "exercicio_organizacao": ["Duas filas", "Meio-campo"],
                "exercicio_coaching_points": ["Finalizar alto", "Conter a bola"],
            },
            follow_redirects=False,
        )
        assert aplicacao.status_code == 303
        assert aplicacao.headers["location"].endswith(f"/treino/sessoes/{sessao.id}")

        db.expire_all()
        importacao = db.query(ImportacaoPlanoTreino).execution_options(ignorar_clube=True).filter(ImportacaoPlanoTreino.id == importacao_id).one()
        sessao_guardada = db.query(SessaoTreino).execution_options(ignorar_clube=True).filter(SessaoTreino.id == sessao.id).one()
        assert importacao.estado == "aplicada"
        assert [(item.nome, item.categoria, item.minutos) for item in importacao.plano.exercicios] == [
            ("Finalização", "TIO", 15),
            ("Defesa 1x1", "TID", 20),
        ]
        assert sessao_guardada.min_tecnica_individual_ofensiva == 15
        assert sessao_guardada.min_tecnica_individual_defensiva == 20
        assert sessao_guardada.volume_minutos == 35

        repeticao = client_admin.post(
            f"/treino/importacoes/{importacao_id}/aplicar",
            data={"data": "2026-08-24"},
            follow_redirects=False,
        )
        assert repeticao.status_code == 303
        db.expire_all()
        plano_guardado = db.query(PlanoTreino).execution_options(ignorar_clube=True).filter(PlanoTreino.id == plano.id).one()
        assert len(plano_guardado.exercicios) == 2

        backup = client_admin.get(f"/exportacoes/equipas/{equipa_id}/dados")
        assert backup.status_code == 200
        with ZipFile(BytesIO(backup.content)) as arquivo:
            nomes_backup = arquivo.namelist()
            assert any(nome.endswith(f"{importacao_id}_plano-teste.docx") for nome in nomes_backup)
            assert any(nome.endswith("treinos_e_planos.xlsx") for nome in nomes_backup)
    finally:
        if importacao_id:
            importacao = db.query(ImportacaoPlanoTreino).execution_options(ignorar_clube=True).filter(ImportacaoPlanoTreino.id == importacao_id).first()
            if importacao:
                db.delete(importacao)
                db.commit()
        if macro_id:
            macro = db.query(Macrociclo).execution_options(ignorar_clube=True).filter(Macrociclo.id == macro_id).first()
            if macro:
                db.delete(macro)
                db.commit()
        if equipa_id:
            equipa = db.query(Equipa).execution_options(ignorar_clube=True).filter(Equipa.id == equipa_id).first()
            if equipa:
                db.delete(equipa)
                db.commit()
        db.close()
