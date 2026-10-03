"""Backups editáveis e isolados por clube e equipa."""

import json
import re
import unicodedata
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from fastapi import APIRouter, Depends, Form, HTTPException, Request
from fastapi.responses import HTMLResponse, RedirectResponse, Response, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pypdf import PdfReader, PdfWriter
from sqlalchemy import and_, or_, select
from sqlalchemy.inspection import inspect as sa_inspect
from sqlalchemy.orm import Session, selectinload
from template_config import Jinja2Templates

from auth import verificar_acesso_equipa
from auth_deps import exigir_autenticacao, exigir_coordenacao, exigir_operacao_desportiva, verificar_ou_negar_equipa
from database import get_db
from flash import set_flash
from models.avaliacao import AvaliacaoJogador, ParametroAvaliacao
from models.assiduidade import RegistoAssiduidade
from models.atividade import Atividade
from models.atribuicao_espaco_semanal import AtribuicaoEspacoSemanal
from models.auditoria import RegistoAuditoria
from models.documento_jogador import DocumentoJogador
from models.equipa import Equipa
from models.espaco_agenda import AgendaEspacoEquipa
from models.espaco_treino import DisponibilidadeEspacoTreino, EspacoTreino, IndisponibilidadeEspacoTreino
from models.fisico import PlaneamentoFisico, TesteFisico
from models.identidade_clube import IdentidadeClube
from models.indisponibilidade_treinador import IndisponibilidadeTreinador
from models.historico_planeamento import HistoricoPlaneamentoClube
from models.jogador import Jogador, RegistoAntropometrico
from models.jogo import Convocado, Jogo, PlanoJogo
from models.notificacao import Notificacao
from models.planeamento import (
    ConteudoSessaoTreino,
    Macrociclo,
    Mesociclo,
    Microciclo,
    PlaneamentoAnualEquipa,
    SessaoTreino,
)
from models.planeamento_clube import (
    EventoPlaneamentoClube,
    PlaneamentoClubeAnual,
    PlaneamentoClubeMensal,
    PlanoSemanalClube,
    evento_planeamento_clube_equipa,
)
from models.playbook import PassoSetPlay, Playbook, SetPlay
from models.relatorio_exportacao import PredefinicaoRelatorio
from models.scouting import (
    ScoutingAdversario,
    ScoutingAdversarioItem,
    ScoutingAdversarioJogo,
    ScoutingBlocoColetivo,
    ScoutingConceitoOfensivo,
    ScoutingItemBiblioteca,
    ScoutingJogador,
    ScoutingJogadorItem,
    ScoutingLiderEstatistico,
    ScoutingMomentoPlaybook,
    ScoutingRotacaoPrevista,
    ScoutingSubstituicaoPrevista,
    STAT_FIELDS,
)
from models.treinador import Treinador, treinador_equipa
from models.treino import BibliotecoExercicio, BlocoTreino, ExercicioPlano, ImportacaoPlanoTreino, PlanoTreino
from models.utilizador import Utilizador
from paths import STATIC_DIR, TEMPLATES_DIR
from services.auditoria_service import registar_auditoria
from services.fotografia_jogador_service import caminho_fotografia
from services.identidade_clube_service import caminho_ficheiro_logo
from services.desenho_service import (
    compor_desenho_sobre_campo,
    decompor_desenho_editavel,
    imagem_estatica_data_uri,
    normalizar_data_uri,
)
from services.pdf_service import content_disposition_pdf, preparar_contexto_pdf, renderizar_pdf
from services.plano_jogo_service import graficos_plano, jogo_em_preparacao, jogos_em_preparacao, rotacoes_por_grupo
from services.scouting_individual_service import construir_perfil_individual

router = APIRouter(prefix="/exportacoes", tags=["exportacoes"])
templates = Jinja2Templates(directory=str(TEMPLATES_DIR))

COLUNAS_SENSIVEIS = {"utilizadores": {"password_hash", "sessao_versao"}}
LIMITE_TEXTO_EXCEL = 32_000


def _nome_seguro(valor: str, fallback: str = "dados") -> str:
    normalizado = unicodedata.normalize("NFKD", valor or "")
    ascii_texto = normalizado.encode("ascii", "ignore").decode("ascii")
    nome = re.sub(r"[^A-Za-z0-9._-]+", "_", ascii_texto).strip("._-")
    return nome or fallback


def _valor_serializavel(valor):
    if valor is None or isinstance(valor, (str, int, float, bool)):
        return valor
    if isinstance(valor, (date, datetime)):
        return valor.isoformat()
    return getattr(valor, "value", str(valor))


def _juntar_pdfs(documentos: list[tuple[str, bytes]]) -> bytes:
    escritor = PdfWriter()
    for _, conteudo in documentos:
        leitor = PdfReader(BytesIO(conteudo))
        for pagina in leitor.pages:
            escritor.add_page(pagina)
    destino = BytesIO()
    escritor.write(destino)
    return destino.getvalue()


def _arquivar_pdfs(documentos: list[tuple[str, bytes]]) -> bytes:
    destino = BytesIO()
    with ZipFile(destino, "w", ZIP_DEFLATED) as arquivo:
        for indice, (nome, conteudo) in enumerate(documentos, start=1):
            arquivo.writestr(f"{indice:02d}_{_nome_seguro(nome, 'documento')}.pdf", conteudo)
    return destino.getvalue()


def _gerar_pdf_planeamento_selecionado(
    tipo: str,
    item_id: int,
    equipa_id: int,
    request: Request,
    db: Session,
    utilizador,
) -> tuple[str, bytes]:
    # As rotas individuais continuam a ser a fonte única dos dados e do design de cada PDF.
    from routers.planeamento import exportar_pdf_macrociclo, exportar_pdf_mesociclo, exportar_pdf_microciclo
    from routers.treino import exportar_pdf_plano

    if tipo == "macro":
        item = db.query(Macrociclo).filter(Macrociclo.id == item_id, Macrociclo.equipa_id == equipa_id).first()
        if not item:
            raise HTTPException(status_code=400, detail="Macrociclo inválido para esta equipa.")
        resposta = exportar_pdf_macrociclo(item.id, request, db, utilizador)
        nome = f"Macrociclo {item.numero}"
    elif tipo == "meso":
        item = db.query(Mesociclo).join(Macrociclo).filter(
            Mesociclo.id == item_id,
            Macrociclo.equipa_id == equipa_id,
        ).first()
        if not item:
            raise HTTPException(status_code=400, detail="Mesociclo inválido para esta equipa.")
        resposta = exportar_pdf_mesociclo(item.id, request, db, utilizador)
        nome = f"Mesociclo {item.numero}"
    elif tipo == "micro":
        item = db.query(Microciclo).join(Mesociclo).join(Macrociclo).filter(
            Microciclo.id == item_id,
            Macrociclo.equipa_id == equipa_id,
        ).first()
        if not item:
            raise HTTPException(status_code=400, detail="Microciclo inválido para esta equipa.")
        resposta = exportar_pdf_microciclo(item.id, request, db, utilizador)
        nome = f"Microciclo {item.numero}"
    elif tipo == "treino":
        item = db.query(SessaoTreino).join(Microciclo).join(Mesociclo).join(Macrociclo).filter(
            SessaoTreino.id == item_id,
            Macrociclo.equipa_id == equipa_id,
            SessaoTreino.ativa.is_(True),
            SessaoTreino.e_jogo.is_(False),
        ).first()
        if not item:
            raise HTTPException(status_code=400, detail="Treino inválido para esta equipa.")
        resposta = exportar_pdf_plano(item.id, request, db, utilizador)
        prefixo = "Treino físico" if (item.tipo_treino or "normal") == "fisico" else "Treino"
        nome = f"{prefixo} {item.numero_treino}" if item.numero_treino is not None else prefixo
    else:
        raise HTTPException(status_code=400, detail="Tipo de documento inválido.")

    if not isinstance(resposta, Response) or resposta.media_type != "application/pdf":
        raise HTTPException(status_code=400, detail=f"Não foi possível gerar {nome}.")
    return nome, bytes(resposta.body)


@router.get("/equipas/{equipa_id}/planeamento", response_class=HTMLResponse)
def selecionar_exportacao_planeamento(
    equipa_id: int,
    request: Request,
    db: Session = Depends(get_db),
    utilizador=Depends(exigir_autenticacao),
):
    verificar_ou_negar_equipa(utilizador, equipa_id, db)
    equipa = db.query(Equipa).filter(Equipa.id == equipa_id).first()
    if not equipa:
        return RedirectResponse("/equipas", status_code=303)

    macros = db.query(Macrociclo).filter(Macrociclo.equipa_id == equipa_id).order_by(
        Macrociclo.numero.asc(), Macrociclo.data_inicio.asc(), Macrociclo.id.asc(),
    ).all()
    mesos = db.query(Mesociclo).join(Macrociclo).filter(Macrociclo.equipa_id == equipa_id).order_by(
        Macrociclo.numero.asc(), Mesociclo.numero.asc(), Mesociclo.data_inicio.asc(), Mesociclo.id.asc(),
    ).all()
    micros = db.query(Microciclo).join(Mesociclo).join(Macrociclo).filter(Macrociclo.equipa_id == equipa_id).order_by(
        Microciclo.numero.asc(), Microciclo.data_inicio.asc(), Microciclo.id.asc(),
    ).all()
    treinos = db.query(SessaoTreino).join(Microciclo).join(Mesociclo).join(Macrociclo).filter(
        Macrociclo.equipa_id == equipa_id,
        SessaoTreino.ativa.is_(True),
        SessaoTreino.e_jogo.is_(False),
    ).order_by(
        SessaoTreino.tipo_treino.asc(), SessaoTreino.numero_treino.asc(),
        SessaoTreino.data.asc(), SessaoTreino.hora_inicio.asc(), SessaoTreino.id.asc(),
    ).all()
    return templates.TemplateResponse(request, "exportacoes/planeamento_selecao.html", {
        "utilizador": utilizador,
        "equipa": equipa,
        "macrociclos": macros,
        "mesociclos": mesos,
        "microciclos": micros,
        "treinos": treinos,
    })


@router.post("/equipas/{equipa_id}/planeamento")
def exportar_planeamento_selecionado(
    equipa_id: int,
    request: Request,
    itens: list[str] = Form([]),
    formato: str = Form("pdf"),
    db: Session = Depends(get_db),
    utilizador=Depends(exigir_autenticacao),
):
    verificar_ou_negar_equipa(utilizador, equipa_id, db)
    equipa = db.query(Equipa).filter(Equipa.id == equipa_id).first()
    if not equipa:
        return RedirectResponse("/equipas", status_code=303)
    if not itens:
        set_flash(request, "Seleciona pelo menos um documento para exportar.", "error")
        return RedirectResponse(f"/exportacoes/equipas/{equipa_id}/planeamento", status_code=303)
    if len(itens) > 80:
        raise HTTPException(status_code=400, detail="Seleciona no máximo 80 documentos de cada vez.")
    if formato not in {"pdf", "zip"}:
        raise HTTPException(status_code=400, detail="Formato de exportação inválido.")

    documentos: list[tuple[str, bytes]] = []
    vistos: set[tuple[str, int]] = set()
    for valor in itens:
        try:
            tipo, id_texto = valor.split(":", 1)
            chave = (tipo, int(id_texto))
        except (AttributeError, TypeError, ValueError):
            raise HTTPException(status_code=400, detail="Seleção de documentos inválida.")
        if chave in vistos:
            continue
        vistos.add(chave)
        documentos.append(_gerar_pdf_planeamento_selecionado(
            chave[0], chave[1], equipa_id, request, db, utilizador,
        ))

    nome_base = f"Planeamento {equipa.nome}"
    if formato == "zip":
        return Response(
            _arquivar_pdfs(documentos),
            media_type="application/zip",
            headers={"Content-Disposition": f'attachment; filename="{_nome_seguro(nome_base)}.zip"'},
        )
    return Response(
        _juntar_pdfs(documentos),
        media_type="application/pdf",
        headers={"Content-Disposition": content_disposition_pdf(nome_base)},
    )


def _valor_excel(valor):
    """Mantém datas editáveis e neutraliza fórmulas introduzidas nos dados."""
    if valor is None or isinstance(valor, (int, float, bool, date, datetime)):
        return valor
    texto = str(getattr(valor, "value", valor))
    if texto.startswith(("=", "+", "-", "@")):
        texto = "'" + texto
    if len(texto) > LIMITE_TEXTO_EXCEL:
        return texto[:LIMITE_TEXTO_EXCEL] + "\n[Conteúdo integral no JSON desta pasta]"
    return texto


def _registo_para_dict(registo, excluir: set[str] | None = None) -> dict:
    if isinstance(registo, dict):
        return {str(chave): _valor_serializavel(valor) for chave, valor in registo.items()}
    excluidas = excluir or set()
    return {
        coluna.key: _valor_serializavel(getattr(registo, coluna.key))
        for coluna in sa_inspect(registo).mapper.column_attrs
        if coluna.key not in excluidas
    }


def _dados_conjunto(nome: str, linhas: list) -> list[dict]:
    return [_registo_para_dict(linha, COLUNAS_SENSIVEIS.get(nome, set())) for linha in linhas]


def _criar_excel(titulo: str, conjuntos: list[tuple[str, list]]) -> bytes:
    livro = Workbook()
    livro.remove(livro.active)
    livro.properties.title = titulo
    livro.properties.creator = "Plataforma de Gestão de Basquetebol"

    for nome, linhas in conjuntos:
        folha = livro.create_sheet(nome[:31])
        dados = _dados_conjunto(nome, linhas)
        if not dados:
            folha.append(["Sem registos"])
            folha["A1"].font = Font(bold=True, color="64748B")
            folha.column_dimensions["A"].width = 22
            continue

        colunas = list(dados[0].keys())
        folha.append(colunas)
        for celula in folha[1]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor="1E293B")
            celula.alignment = Alignment(vertical="center")
        folha.row_dimensions[1].height = 24
        for registo in dados:
            folha.append([_valor_excel(registo.get(coluna)) for coluna in colunas])

        for indice, coluna in enumerate(colunas, start=1):
            maior = max(len(str(coluna)), *(len(str(linha.get(coluna) or "")) for linha in dados))
            folha.column_dimensions[folha.cell(1, indice).column_letter].width = min(max(maior + 2, 12), 48)
        for linha in folha.iter_rows(min_row=2):
            for celula in linha:
                celula.alignment = Alignment(vertical="top", wrap_text=True)
        folha.freeze_panes = "A2"
        folha.auto_filter.ref = folha.dimensions

    resultado = BytesIO()
    livro.save(resultado)
    return resultado.getvalue()


def _adicionar_tema(zipado: ZipFile, pasta: str, titulo: str, conjuntos: list[tuple[str, list]]) -> None:
    nome = _nome_seguro(titulo).lower()
    zipado.writestr(f"{pasta}/{nome}.xlsx", _criar_excel(titulo, conjuntos))
    integral = {nome_folha: _dados_conjunto(nome_folha, linhas) for nome_folha, linhas in conjuntos}
    zipado.writestr(f"{pasta}/{nome}_integral.json", json.dumps(integral, ensure_ascii=False, indent=2))


def _adicionar_fotografias(zipado: ZipFile, jogadores: list[Jogador], prefixo: str) -> None:
    for jogador in jogadores:
        fotografia = caminho_fotografia(jogador.fotografia)
        if fotografia and fotografia.is_file():
            zipado.write(
                fotografia,
                f"{prefixo}/Ficheiros/Fotografias/{jogador.id}_{_nome_seguro(fotografia.name)}",
            )


def _adicionar_planos_importados(
    zipado: ZipFile,
    importacoes: list[ImportacaoPlanoTreino],
    prefixo: str,
) -> None:
    for importacao in importacoes:
        zipado.writestr(
            f"{prefixo}/Ficheiros/Planos_importados/{importacao.id}_{_nome_seguro(importacao.nome_original)}",
            importacao.ficheiro,
        )


def _adicionar_logos(zipado: ZipFile, clube, identidade: IdentidadeClube | None) -> None:
    urls = {
        "logo_principal_clube": clube.logo_principal,
        "logo_secundario_clube": clube.logo_secundario,
        "logo_principal_identidade": getattr(identidade, "logo_principal", None),
        "logo_secundario_identidade": getattr(identidade, "logo_secundario", None),
    }
    adicionados: set[Path] = set()
    for nome, url in urls.items():
        caminho = caminho_ficheiro_logo(STATIC_DIR, url)
        if caminho and caminho.is_file() and caminho not in adicionados:
            adicionados.add(caminho)
            zipado.write(caminho, f"00_Clube/Ficheiros/Logotipos/{nome}{caminho.suffix.lower()}")


def _linhas_associacao(db: Session, tabela, condicao) -> list[dict]:
    return [dict(linha) for linha in db.execute(select(tabela).where(condicao)).mappings().all()]


def _conteudo_equipa(
    db: Session,
    equipa: Equipa,
) -> tuple[list[tuple[str, str, list]], list[Jogador], list[ImportacaoPlanoTreino]]:
    equipa_id = equipa.id
    jogadores = db.query(Jogador).filter(Jogador.equipa_id == equipa_id).order_by(Jogador.nome).all()
    jogadores_ids = [jogador.id for jogador in jogadores]
    jogos = db.query(Jogo).filter(Jogo.equipa_id == equipa_id).order_by(Jogo.data).all()
    jogos_ids = [jogo.id for jogo in jogos]

    anuais = db.query(PlaneamentoAnualEquipa).filter(PlaneamentoAnualEquipa.equipa_id == equipa_id).all()
    macros = db.query(Macrociclo).filter(Macrociclo.equipa_id == equipa_id).order_by(Macrociclo.data_inicio).all()
    macros_ids = [macro.id for macro in macros]
    mesos = db.query(Mesociclo).filter(Mesociclo.macrociclo_id.in_(macros_ids)).order_by(Mesociclo.data_inicio).all() if macros_ids else []
    mesos_ids = [meso.id for meso in mesos]
    micros = db.query(Microciclo).filter(Microciclo.mesociclo_id.in_(mesos_ids)).order_by(Microciclo.data_inicio).all() if mesos_ids else []
    micros_ids = [micro.id for micro in micros]
    sessoes = db.query(SessaoTreino).filter(SessaoTreino.microciclo_id.in_(micros_ids)).order_by(SessaoTreino.data).all() if micros_ids else []
    sessoes_ids = [sessao.id for sessao in sessoes]
    planos = db.query(PlanoTreino).filter(PlanoTreino.sessao_id.in_(sessoes_ids)).all() if sessoes_ids else []
    planos_ids = [plano.id for plano in planos]
    importacoes = db.query(ImportacaoPlanoTreino).filter(
        ImportacaoPlanoTreino.plano_id.in_(planos_ids)
    ).all() if planos_ids else []
    importacoes_metadados = [
        _registo_para_dict(importacao, {"ficheiro"}) for importacao in importacoes
    ]

    relacoes_treinadores = _linhas_associacao(db, treinador_equipa, treinador_equipa.c.equipa_id == equipa_id)
    treinadores_ids = [linha["treinador_id"] for linha in relacoes_treinadores]
    treinadores = db.query(Treinador).filter(Treinador.id.in_(treinadores_ids)).order_by(Treinador.nome).all() if treinadores_ids else []
    utilizadores = db.query(Utilizador).filter(Utilizador.treinador_id.in_(treinadores_ids)).all() if treinadores_ids else []

    eventos_associados = select(evento_planeamento_clube_equipa.c.evento_id).where(
        evento_planeamento_clube_equipa.c.equipa_id == equipa_id
    )
    eventos = db.query(EventoPlaneamentoClube).filter(
        EventoPlaneamentoClube.clube_id == equipa.clube_id,
        or_(
            EventoPlaneamentoClube.equipa_id == equipa_id,
            EventoPlaneamentoClube.aplica_a_todas_equipas.is_(True),
            EventoPlaneamentoClube.id.in_(eventos_associados),
        ),
    ).order_by(EventoPlaneamentoClube.data).all()
    eventos_ids = [evento.id for evento in eventos]
    relacoes_eventos = _linhas_associacao(
        db, evento_planeamento_clube_equipa, evento_planeamento_clube_equipa.c.evento_id.in_(eventos_ids)
    ) if eventos_ids else []
    atividades_ids = [evento.atividade_id for evento in eventos if evento.atividade_id]

    playbooks = db.query(Playbook).filter(Playbook.equipa_id == equipa_id).all()
    playbooks_ids = [playbook.id for playbook in playbooks]
    setplays = db.query(SetPlay).filter(SetPlay.playbook_id.in_(playbooks_ids)).all() if playbooks_ids else []
    setplays_ids = [setplay.id for setplay in setplays]
    agenda_espacos = db.query(AgendaEspacoEquipa).filter(AgendaEspacoEquipa.equipa_id == equipa_id).all()
    atribuicoes_espacos = db.query(AtribuicaoEspacoSemanal).filter(
        AtribuicaoEspacoSemanal.equipa_id == equipa_id
    ).all()
    espacos_ids = {
        item.espaco_id for item in [*agenda_espacos, *atribuicoes_espacos] if item.espaco_id
    }
    espacos = db.query(EspacoTreino).filter(EspacoTreino.id.in_(espacos_ids)).all() if espacos_ids else []

    temas = [
        ("00_Equipa", "equipa_e_equipa_tecnica", [
            ("equipa", [equipa]), ("treinadores", treinadores),
            ("treinador_equipa", relacoes_treinadores), ("utilizadores", utilizadores),
        ]),
        ("01_Plantel", "plantel", [
            ("jogadores", jogadores),
            ("documentos_jogadores", db.query(DocumentoJogador).filter(DocumentoJogador.jogador_id.in_(jogadores_ids)).all() if jogadores_ids else []),
            ("registos_antropometricos", db.query(RegistoAntropometrico).filter(RegistoAntropometrico.jogador_id.in_(jogadores_ids)).all() if jogadores_ids else []),
            ("testes_fisicos", db.query(TesteFisico).filter(TesteFisico.jogador_id.in_(jogadores_ids)).all() if jogadores_ids else []),
        ]),
        ("02_Planeamento", "planeamento_da_equipa", [
            ("planeamentos_anuais_equipa", anuais), ("macrociclos", macros),
            ("mesociclos", mesos), ("microciclos", micros),
            ("eventos_planeamento_clube", eventos), ("evento_planeamento_equipas", relacoes_eventos),
            ("atividades", db.query(Atividade).filter(Atividade.id.in_(atividades_ids)).all() if atividades_ids else []),
        ]),
        ("03_Treinos", "treinos_e_planos", [
            ("sessoes_treino", sessoes),
            ("conteudos_sessao_treino", db.query(ConteudoSessaoTreino).filter(ConteudoSessaoTreino.sessao_id.in_(sessoes_ids)).all() if sessoes_ids else []),
            ("planos_treino", planos),
            ("importacoes_plano_treino", importacoes_metadados),
            ("blocos_treino", db.query(BlocoTreino).filter(BlocoTreino.plano_id.in_(planos_ids)).all() if planos_ids else []),
            ("exercicios_plano", db.query(ExercicioPlano).filter(ExercicioPlano.plano_id.in_(planos_ids)).all() if planos_ids else []),
            ("registos_assiduidade", db.query(RegistoAssiduidade).filter(RegistoAssiduidade.sessao_id.in_(sessoes_ids)).all() if sessoes_ids else []),
        ]),
        ("04_Jogos", "jogos_convocatorias_e_planos", [
            ("jogos", jogos),
            ("convocados", db.query(Convocado).filter(Convocado.jogo_id.in_(jogos_ids)).all() if jogos_ids else []),
            ("planos_jogo", db.query(PlanoJogo).filter(PlanoJogo.jogo_id.in_(jogos_ids)).all() if jogos_ids else []),
        ]),
        ("05_Preparacao_fisica", "preparacao_fisica", [
            ("planeamentos_fisicos", db.query(PlaneamentoFisico).filter(PlaneamentoFisico.equipa_id == equipa_id).all()),
            ("parametros_avaliacao", db.query(ParametroAvaliacao).filter(ParametroAvaliacao.equipa_id == equipa_id).all()),
            ("avaliacoes_jogadores", db.query(AvaliacaoJogador).filter(AvaliacaoJogador.equipa_id == equipa_id).all()),
        ]),
        ("06_Playbook", "playbook", [
            ("playbooks", playbooks), ("setplays", setplays),
            ("passos_setplay", db.query(PassoSetPlay).filter(PassoSetPlay.setplay_id.in_(setplays_ids)).all() if setplays_ids else []),
        ]),
        ("07_Agenda_e_espacos", "agenda_e_espacos", [
            ("espacos_treino", espacos),
            ("agenda_espaco_equipa", agenda_espacos),
            ("atribuicoes_espaco_semanais", atribuicoes_espacos),
        ]),
        ("08_Biblioteca", "biblioteca_de_exercicios_do_clube", [
            ("biblioteca_exercicios", db.query(BibliotecoExercicio).filter(BibliotecoExercicio.clube_id == equipa.clube_id).all()),
        ]),
    ]
    return temas, jogadores, importacoes


def _escrever_equipa(zipado: ZipFile, db: Session, equipa: Equipa, raiz: str) -> dict:
    temas, jogadores, importacoes = _conteudo_equipa(db, equipa)
    contagens = {}
    for pasta, titulo, conjuntos in temas:
        _adicionar_tema(zipado, f"{raiz}/{pasta}", titulo, conjuntos)
        contagens[pasta] = sum(len(linhas) for _, linhas in conjuntos)
    _adicionar_fotografias(zipado, jogadores, f"{raiz}/01_Plantel")
    _adicionar_planos_importados(zipado, importacoes, f"{raiz}/03_Treinos")
    return contagens


def _escrever_manifesto(zipado: ZipFile, clube, utilizador, equipas: list[dict], escopo: str) -> None:
    manifesto = {
        "formato_backup": "basquetebol-editavel-v2",
        "criado_em_utc": datetime.now(timezone.utc).isoformat(),
        "escopo": escopo,
        "clube": {"id": clube.id, "nome_oficial": clube.nome_oficial, "nome_curto": clube.nome_curto},
        "exportado_por": {"id": utilizador.id, "username": utilizador.username, "perfil": _valor_serializavel(utilizador.perfil)},
        "equipas": equipas,
        "seguranca": "Hashes de palavra-passe, versões de sessão e tokens de recuperação não são exportados.",
    }
    zipado.writestr("MANIFESTO.json", json.dumps(manifesto, ensure_ascii=False, indent=2))
    zipado.writestr(
        "LEIA-ME.txt",
        "BACKUP EDITÁVEL DA PLATAFORMA DE BASQUETEBOL\n\n"
        "Os dados estão separados por tema. Cada pasta contém um Excel para consulta e edição "
        "e um JSON integral que preserva conteúdos extensos, incluindo desenhos táticos.\n"
        "Alterar estes ficheiros não altera automaticamente a plataforma. Este pacote não inclui "
        "palavras-passe, tokens de recuperação ou chaves de sessão.\n",
    )


def _historico_do_clube(
    db: Session,
    anuais: list[PlaneamentoClubeAnual],
    mensais: list[PlaneamentoClubeMensal],
    semanais: list[PlanoSemanalClube],
) -> list[HistoricoPlaneamentoClube]:
    """Filtra a tabela histórica antiga, que ainda não tem clube_id."""
    grupos = {
        "planeamento_anual": [item.id for item in anuais],
        "planeamento_mensal": [item.id for item in mensais],
        "plano_semanal": [item.id for item in semanais],
    }
    filtros = [
        and_(
            HistoricoPlaneamentoClube.entidade == entidade,
            HistoricoPlaneamentoClube.entidade_id.in_(ids),
        )
        for entidade, ids in grupos.items()
        if ids
    ]
    if not filtros:
        return []
    return db.query(HistoricoPlaneamentoClube).filter(or_(*filtros)).order_by(
        HistoricoPlaneamentoClube.criado_em
    ).all()


@router.get("/equipas/{equipa_id}/dados")
def exportar_dados_equipa(
    equipa_id: int,
    request: Request,
    db: Session = Depends(get_db),
    utilizador=Depends(exigir_operacao_desportiva),
):
    """Descarrega um ZIP editável da equipa, limitado pelas permissões atuais."""
    verificar_ou_negar_equipa(utilizador, equipa_id, db)
    equipa = db.query(Equipa).filter(Equipa.id == equipa_id).first()
    if equipa is None:
        raise HTTPException(status_code=404, detail="Equipa não encontrada.")

    clube = getattr(request.state, "clube_atual", None)
    if clube is None or equipa.clube_id != clube.id:
        raise HTTPException(status_code=403, detail="A equipa não pertence ao clube selecionado.")

    ficheiro = BytesIO()
    raiz = f"Equipa_{equipa.id}_{_nome_seguro(equipa.nome)}"
    with ZipFile(ficheiro, "w", ZIP_DEFLATED) as zipado:
        identidade = db.query(IdentidadeClube).filter(IdentidadeClube.clube_id == clube.id).first()
        _adicionar_tema(
            zipado,
            "00_Clube",
            "identificacao_do_clube",
            [("clube", [clube]), ("identidade_clube", [identidade] if identidade else [])],
        )
        _adicionar_logos(zipado, clube, identidade)
        contagens = _escrever_equipa(zipado, db, equipa, raiz)
        _escrever_manifesto(
            zipado, clube, utilizador,
            [{"id": equipa.id, "nome": equipa.nome, "pasta": raiz, "contagens": contagens}],
            "equipa",
        )

    registar_auditoria(
        db, utilizador, "exportou", "backup_equipa", equipa_id,
        f"Backup editável completo da equipa {equipa.nome}.", request,
    )
    db.commit()
    ficheiro.seek(0)
    return StreamingResponse(
        ficheiro,
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="backup_equipa_{_nome_seguro(equipa.nome)}.zip"'},
    )


@router.get("/clube/dados")
def exportar_dados_clube(
    request: Request,
    db: Session = Depends(get_db),
    utilizador=Depends(exigir_coordenacao),
):
    """Descarrega o clube por temas, com uma pasta detalhada por equipa."""
    clube = getattr(request.state, "clube_atual", None)
    if clube is None:
        raise HTTPException(status_code=403, detail="Seleciona um clube antes de exportar dados.")

    identidade = db.query(IdentidadeClube).filter(IdentidadeClube.clube_id == clube.id).first()
    equipas = db.query(Equipa).filter(Equipa.clube_id == clube.id).order_by(Equipa.nome).all()
    utilizadores = db.query(Utilizador).filter(Utilizador.clube_id == clube.id).order_by(Utilizador.nome).all()
    treinadores = db.query(Treinador).filter(Treinador.clube_id == clube.id).order_by(Treinador.nome).all()
    treinadores_ids = [treinador.id for treinador in treinadores]
    equipa_ids = [equipa.id for equipa in equipas]

    anuais = db.query(PlaneamentoClubeAnual).filter(PlaneamentoClubeAnual.clube_id == clube.id).all()
    mensais = db.query(PlaneamentoClubeMensal).filter(PlaneamentoClubeMensal.clube_id == clube.id).all()
    semanais = db.query(PlanoSemanalClube).filter(PlanoSemanalClube.clube_id == clube.id).all()
    eventos = db.query(EventoPlaneamentoClube).filter(EventoPlaneamentoClube.clube_id == clube.id).all()
    eventos_ids = [evento.id for evento in eventos]

    temas_gerais = [
        ("00_Clube", "identificacao_do_clube", [
            ("clube", [clube]), ("identidade_clube", [identidade] if identidade else []), ("equipas", equipas),
        ]),
        ("01_Utilizadores", "utilizadores_e_equipa_tecnica", [
            ("utilizadores", utilizadores), ("treinadores", treinadores),
            ("treinador_equipa", _linhas_associacao(db, treinador_equipa, treinador_equipa.c.equipa_id.in_(equipa_ids)) if equipa_ids else []),
            ("indisponibilidades_treinador", db.query(IndisponibilidadeTreinador).filter(IndisponibilidadeTreinador.treinador_id.in_(treinadores_ids)).all() if treinadores_ids else []),
        ]),
        ("02_Planeamento_do_clube", "planeamento_do_clube", [
            ("planeamentos_clube_anuais", anuais), ("planeamentos_clube_mensais", mensais),
            ("planos_semanais_clube", semanais), ("eventos_planeamento_clube", eventos),
            ("evento_planeamento_equipas", _linhas_associacao(db, evento_planeamento_clube_equipa, evento_planeamento_clube_equipa.c.evento_id.in_(eventos_ids)) if eventos_ids else []),
            ("atividades", db.query(Atividade).filter(Atividade.clube_id == clube.id).all()),
        ]),
        ("03_Espacos_e_horarios", "espacos_horarios_e_atribuicoes", [
            ("espacos_treino", db.query(EspacoTreino).filter(EspacoTreino.clube_id == clube.id).all()),
            ("disponibilidades_espaco_treino", db.query(DisponibilidadeEspacoTreino).filter(DisponibilidadeEspacoTreino.clube_id == clube.id).all()),
            ("indisponibilidades_espaco_treino", db.query(IndisponibilidadeEspacoTreino).filter(IndisponibilidadeEspacoTreino.clube_id == clube.id).all()),
            ("atribuicoes_espaco_semanais", db.query(AtribuicaoEspacoSemanal).filter(AtribuicaoEspacoSemanal.clube_id == clube.id).all()),
            ("agenda_espaco_equipa", db.query(AgendaEspacoEquipa).filter(AgendaEspacoEquipa.clube_id == clube.id).all()),
        ]),
        ("04_Biblioteca", "biblioteca_de_exercicios", [
            ("biblioteca_exercicios", db.query(BibliotecoExercicio).filter(BibliotecoExercicio.clube_id == clube.id).all()),
        ]),
        ("05_Auditoria_e_notificacoes", "auditoria_e_notificacoes", [
            ("registos_auditoria", db.query(RegistoAuditoria).filter(RegistoAuditoria.clube_id == clube.id).all()),
            ("historico_planeamento_clube", _historico_do_clube(db, anuais, mensais, semanais)),
            ("notificacoes", db.query(Notificacao).filter(Notificacao.clube_id == clube.id).all()),
        ]),
    ]

    ficheiro = BytesIO()
    manifesto_equipas = []
    with ZipFile(ficheiro, "w", ZIP_DEFLATED) as zipado:
        for pasta, titulo, conjuntos in temas_gerais:
            _adicionar_tema(zipado, pasta, titulo, conjuntos)
        _adicionar_logos(zipado, clube, identidade)
        for equipa in equipas:
            raiz = f"06_Equipas/Equipa_{equipa.id}_{_nome_seguro(equipa.nome)}"
            contagens = _escrever_equipa(zipado, db, equipa, raiz)
            manifesto_equipas.append({"id": equipa.id, "nome": equipa.nome, "pasta": raiz, "contagens": contagens})
        _escrever_manifesto(zipado, clube, utilizador, manifesto_equipas, "clube")

    registar_auditoria(
        db, utilizador, "exportou", "backup_clube", clube.id,
        f"Backup editável completo do clube {clube.nome_curto}.", request,
    )
    db.commit()
    ficheiro.seek(0)
    return StreamingResponse(
        ficheiro,
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="backup_clube_{_nome_seguro(clube.nome_curto)}.zip"'},
    )


# ---------------------------------------------------------------------------
# Relatorios desportivos configuraveis
# ---------------------------------------------------------------------------

FORMATOS_RELATORIO = {
    "completo": {
        "nome": "Completo",
        "descricao": "Dossier detalhado para preparar, arquivar e partilhar com a equipa tecnica.",
    },
    "facil": {
        "nome": "Leitura facil",
        "descricao": "Resumo visual, direto e confortavel para consultar no telemovel ou no balneario.",
    },
    "a3": {
        "nome": "Poster A3",
        "descricao": "Poster editorial de alto impacto para balneario, com fotografias e mensagens essenciais.",
    },
}


def _secao(chave: str, nome: str, descricao: str, formatos=("completo", "facil")) -> dict:
    return {"chave": chave, "nome": nome, "descricao": descricao, "formatos": formatos}


RELATORIOS_CONFIG = {
    "roster": {
        "nome": "Relatorio de plantel",
        "entidades": "jogadores",
        "secoes": [
            _secao("fotografia", "Fotografia", "Identificacao visual dos atletas.", ("completo", "facil", "a3")),
            _secao("identificacao", "Identificacao", "Nome, numero, posicao e idade.", ("completo", "facil", "a3")),
            _secao("fisico", "Perfil fisico", "Altura, peso e mao dominante.", ("completo", "facil")),
            _secao("estado", "Estado desportivo", "Estado atual no roster adversario.", ("completo",)),
            _secao("notas", "Notas tecnicas", "Observacoes e equipa anterior.", ("completo", "facil")),
            _secao("estatisticas", "Estatisticas", "Medias da epoca e indicadores principais.", ("completo", "facil")),
        ],
    },
    "scouting-individual": {
        "nome": "Scouting individual",
        "entidades": "atletas do adversario",
        "secoes": [
            _secao("identificacao", "Identificacao", "Numero, posicao, idade e perfil fisico.", ("completo", "facil", "a3")),
            _secao("fotografia", "Fotografia", "Referencia visual do atleta.", ("completo", "facil", "a3")),
            _secao("pontos_fortes", "Pontos fortes", "Caracteristicas que exigem maior atencao.", ("completo", "facil", "a3")),
            _secao("pontos_fracos", "Pontos a explorar", "Limitacoes e oportunidades identificadas.", ("completo", "facil", "a3")),
            _secao("avaliacoes", "Avaliações 1–5", "Perfil por grupos e características configuráveis.", ("completo", "facil", "a3")),
            _secao("tendencias", "Tendencias", "Movimentos, decisoes e comportamentos recorrentes.", ("completo", "facil")),
            _secao("game_plan", "Game Plan", "Instruções operacionais para enfrentar o jogador.", ("completo", "facil", "a3")),
            _secao("zonas", "Zonas de acao", "Mapa visual das areas preferenciais.", ("completo", "facil")),
            _secao("estatisticas", "Estatisticas", "Medias da epoca e indicadores principais.", ("completo", "facil")),
            _secao("observacoes", "Observacoes", "Resumo tecnico individual em pontos essenciais.", ("completo", "facil", "a3")),
            _secao("recomendacoes", "Recomendacoes", "Orientacoes praticas para o matchup.", ("completo", "facil", "a3")),
        ],
    },
    "scouting-coletivo": {
        "nome": "Scouting coletivo",
        "secoes": [
            _secao("identidade", "Cabeçalho / identificação", "Equipa, jogo e ideia-chave.", ("completo", "facil", "a3")),
            _secao("chaves", "Chaves da equipa", "Mensagens-chave selecionadas pelo treinador.", ("completo", "facil", "a3")),
            _secao("cinco_rotacao", "Cinco e rotações", "Cinco inicial previsto e principais rotações.", ("completo", "facil", "a3")),
            _secao("estatisticas", "Estatísticas da equipa", "Apenas os indicadores selecionados.", ("completo", "facil")),
            _secao("perfil", "Perfil coletivo", "Como a equipa joga, com avaliação simples.", ("completo", "facil", "a3")),
            _secao("analise", "Análise técnico-tática", "Pontos fortes e fragilidades no ataque e na defesa.", ("completo",)),
            _secao("regras_coberturas", "Regras defensivas mais utilizadas", "Opções defensivas usadas pelo adversário.", ("completo", "facil", "a3")),
            _secao("jogadores_chave", "Jogadores mais influentes", "Lideres e prioridades individuais.", ("completo", "facil", "a3")),
            _secao("resposta", "Nossa resposta", "Instruções para contrariar o adversário.", ("completo", "facil", "a3")),
            _secao("prioridades", "Prioridades", "Prioridades ofensivas e defensivas.", ("completo", "facil", "a3")),
            _secao("situacoes", "Transição, ressalto e especiais", "Situações específicas destacadas.", ("completo", "facil")),
            _secao("diagramas", "Ações principais", "Conceitos ofensivos selecionados, apresentados no final.", ("completo", "facil", "a3")),
        ],
    },
    "plano-jogo": {
        "nome": "Plano de jogo",
        "secoes": [
            _secao("dados_jogo", "Dados do jogo", "Data, hora, local, competicao e adversario.", ("completo", "facil", "a3")),
            _secao("convocatoria", "Convocatoria", "Atletas convocados e informacao de encontro.", ("completo", "facil")),
            _secao("cinco_previsto", "Cinco e rotacoes", "Cinco inicial da equipa e previsao do adversario.", ("completo", "facil", "a3")),
            _secao("plano_ofensivo", "Plano ofensivo", "Chaves e movimentos prioritarios.", ("completo", "facil", "a3")),
            _secao("plano_defensivo", "Plano defensivo", "Coberturas e respostas por acao.", ("completo", "facil", "a3")),
            _secao("situacoes_especiais", "Situacoes especiais", "Decisoes para momentos especificos.", ("completo", "facil")),
            _secao("scouting", "Sintese do scouting", "Informacao coletiva importada do adversario.", ("completo", "facil")),
            _secao("mensagem", "Mensagem final", "Rotações, prioridades e mensagem curta para a equipa.", ("completo", "facil", "a3")),
        ],
    },
}

STAT_LABELS = {
    "minutos": "MIN",
    "pontos": "PTS",
    "fg_c": "FG convertidos",
    "fg_t": "FG tentados",
    "dois_c": "2P convertidos",
    "dois_t": "2P tentados",
    "tres_c": "3P convertidos",
    "tres_t": "3P tentados",
    "ll_c": "LL convertidos",
    "ll_t": "LL tentados",
    "ress_of": "Ressaltos ofensivos",
    "ress_def": "Ressaltos defensivos",
    "ress_total": "Ressaltos",
    "assistencias": "Assistências",
    "faltas": "Faltas",
    "turnovers": "Turnovers",
    "roubos": "Roubos",
    "desarmes": "Desarmes",
    "mais_menos": "+/-",
    "eficiencia": "Eficiência",
}

# Os nomes completos mantêm-se na interface de configuração. Nos PDFs usam-se
# códigos compactos para preservar a leitura das tabelas com muitos indicadores.
STAT_CODES = {
    "minutos": "MIN",
    "pontos": "PTS",
    "fg_c": "FGM",
    "fg_t": "FGA",
    "dois_c": "2PM",
    "dois_t": "2PA",
    "tres_c": "3PM",
    "tres_t": "3PA",
    "ll_c": "FTM",
    "ll_t": "FTA",
    "ress_of": "OREB",
    "ress_def": "DREB",
    "ress_total": "REB",
    "assistencias": "AST",
    "faltas": "PF",
    "turnovers": "TOV",
    "roubos": "STL",
    "desarmes": "BLK",
    "mais_menos": "+/-",
    "eficiencia": "EFF",
}

STAT_HIGHLIGHT_OPTIONS = [
    (STAT_CODES[campo], f"{STAT_LABELS[campo]} ({STAT_CODES[campo]})")
    for campo in STAT_FIELDS
] + [(codigo, codigo) for codigo in ("FG%", "2P%", "3P%", "LL%")]
STAT_HIGHLIGHT_CODES = {codigo for codigo, _ in STAT_HIGHLIGHT_OPTIONS}

PERFIS_CONTACTOS = {"superadministrador", "administrador", "coordenador", "secretaria"}


def _perfil(utilizador) -> str:
    return getattr(getattr(utilizador, "perfil", None), "value", str(getattr(utilizador, "perfil", "")))


def _secoes_padrao(tipo: str, formato: str) -> list[str]:
    return [s["chave"] for s in RELATORIOS_CONFIG[tipo]["secoes"] if formato in s["formatos"]]


def _normalizar_secoes_coletivas(secoes: list[str]) -> list[str]:
    equivalencias = {"ataque": ("analise",), "defesa": ("analise", "regras_coberturas")}
    resultado = []
    for secao in secoes:
        for atual in equivalencias.get(secao, (secao,)):
            if atual not in resultado:
                resultado.append(atual)
    return resultado


def _configuracao_relatorio(tipo: str, dados, utilizador) -> tuple[str, str, list[str]]:
    formato = dados.get("formato", "completo")
    if formato not in FORMATOS_RELATORIO:
        formato = "completo"
    orientacao = "paisagem" if formato == "a3" else dados.get("orientacao", "retrato")
    if orientacao not in {"retrato", "paisagem"}:
        orientacao = "retrato"
    validas = {s["chave"] for s in RELATORIOS_CONFIG[tipo]["secoes"]}
    recebidas = dados.getlist("secoes")
    if tipo == "scouting-coletivo":
        recebidas = _normalizar_secoes_coletivas(recebidas)
    escolhidas = [s for s in recebidas if s in validas]
    ordem_recebida = dados.get("ordem", "").split(",")
    if tipo == "scouting-coletivo":
        ordem_recebida = _normalizar_secoes_coletivas(ordem_recebida)
    ordem = [s for s in ordem_recebida if s in escolhidas]
    secoes = ordem + [s for s in escolhidas if s not in ordem]
    if not secoes and not (tipo == "scouting-coletivo" and dados.get("secoes_configuradas") == "1"):
        secoes = _secoes_padrao(tipo, formato)
    if tipo in {"roster", "scouting-individual"} and "fotografia" not in secoes:
        secoes.insert(0, "fotografia")
    if formato == "a3" or _perfil(utilizador) not in PERFIS_CONTACTOS:
        secoes = [s for s in secoes if s != "contactos"]
    return formato, orientacao, secoes


def _configuracao_individual_por_atleta(dados, atletas, secoes_globais: list[str]) -> tuple[dict[int, set[str]], dict[int, list[str]]]:
    """Resolve as escolhas individuais sem alterar o comportamento dos restantes relatórios."""
    validas = {s["chave"] for s in RELATORIOS_CONFIG["scouting-individual"]["secoes"]}
    configurados = {int(valor) for valor in dados.getlist("configurado_jogador") if str(valor).isdigit()}
    ordem_base = [s for s in dados.get("ordem", "").split(",") if s in validas]
    secoes_por_atleta: dict[int, set[str]] = {}
    estatisticas_por_atleta: dict[int, list[str]] = {}
    for atleta in atletas:
        if atleta.id in configurados:
            escolhidas = [s for s in dados.getlist(f"secoes_jogador_{atleta.id}") if s in validas]
            ordenadas = [s for s in ordem_base if s in escolhidas]
            ordenadas.extend(s for s in escolhidas if s not in ordenadas)
            estatisticas = [s for s in dados.getlist(f"estatisticas_jogador_{atleta.id}") if s in STAT_FIELDS]
        else:
            ordenadas = list(secoes_globais)
            estatisticas = [s for s in dados.getlist("estatisticas") if s in STAT_FIELDS]
        secoes_por_atleta[atleta.id] = set(ordenadas)
        estatisticas_por_atleta[atleta.id] = estatisticas
    return secoes_por_atleta, estatisticas_por_atleta


def _fotografia_embebida(url: str | None) -> str | None:
    caminho = caminho_fotografia(url)
    if not caminho or not caminho.is_file():
        return None
    mime = "image/png" if caminho.suffix.lower() == ".png" else "image/jpeg"
    return imagem_estatica_data_uri(caminho, mime)


def _passos_setplay_pdf(setplay) -> list[dict]:
    """Prepara imagens autónomas do Playbook para o HTML/PDF."""
    campo = STATIC_DIR / "campos" / "campo_meio.png"
    passos = []
    for passo in getattr(setplay, "passos", []) or []:
        desenho = decompor_desenho_editavel(passo.desenho)[0] if passo.desenho else None
        imagem_coach = normalizar_data_uri(passo.imagem_coach) if passo.imagem_coach else None
        if desenho and campo.is_file():
            imagem = compor_desenho_sobre_campo(campo, desenho)
        elif imagem_coach:
            imagem = imagem_coach
        else:
            imagem = imagem_estatica_data_uri(campo) if campo.is_file() else None
        passos.append({
            "ordem": (passo.ordem or 0) + 1,
            "observacao": passo.observacao,
            "imagem": imagem,
        })
    return passos


def _idade(data_nascimento) -> int | None:
    if not data_nascimento:
        return None
    hoje = date.today()
    return hoje.year - data_nascimento.year - ((hoje.month, hoje.day) < (data_nascimento.month, data_nascimento.day))


def _alvo_relatorio(tipo: str, alvo_id: int, db: Session, utilizador):
    clube_id = utilizador.clube_id
    if tipo in {"roster", "scouting-individual", "scouting-coletivo"}:
        consulta = db.query(ScoutingAdversario)
        if tipo in {"roster", "scouting-individual"}:
            consulta = consulta.options(
                selectinload(ScoutingAdversario.jogadores)
                .selectinload(ScoutingJogador.itens)
                .joinedload(ScoutingJogadorItem.item)
                .joinedload(ScoutingItemBiblioteca.grupo),
                selectinload(ScoutingAdversario.jogadores).selectinload(ScoutingJogador.estatisticas),
                selectinload(ScoutingAdversario.jogadores).selectinload(ScoutingJogador.zonas),
                selectinload(ScoutingAdversario.jogadores).selectinload(ScoutingJogador.media_epoca),
                selectinload(ScoutingAdversario.jogadores).selectinload(ScoutingJogador.ficha),
            )
        elif tipo == "scouting-coletivo":
            consulta = consulta.options(
                selectinload(ScoutingAdversario.blocos_coletivos),
                selectinload(ScoutingAdversario.rotacao_prevista).joinedload(ScoutingRotacaoPrevista.jogador),
                selectinload(ScoutingAdversario.substituicoes_previstas).joinedload(ScoutingSubstituicaoPrevista.jogador),
                selectinload(ScoutingAdversario.itens_coletivos)
                .joinedload(ScoutingAdversarioItem.item),
                selectinload(ScoutingAdversario.lideres).joinedload(ScoutingLiderEstatistico.jogador),
                selectinload(ScoutingAdversario.conceitos_ofensivos).joinedload(ScoutingConceitoOfensivo.setplay),
                selectinload(ScoutingAdversario.momentos_playbook).joinedload(ScoutingMomentoPlaybook.setplay),
            )
        adversario = consulta.filter(ScoutingAdversario.id == alvo_id, ScoutingAdversario.clube_id == clube_id).first()
        if not adversario:
            raise HTTPException(404, "Adversario nao encontrado.")
        jogos = [
            ligacao.jogo for ligacao in adversario.jogos
            if ligacao.jogo and verificar_acesso_equipa(utilizador, ligacao.jogo.equipa_id, db)
        ]
        if not jogos:
            raise HTTPException(403, "Sem acesso a este adversario.")
        futuros = jogos_em_preparacao(jogos)
        if tipo in {"scouting-individual", "scouting-coletivo"} and not futuros:
            raise HTTPException(409, "Nao existem jogos futuros para preparar este scouting.")
        jogo_referencia = futuros[0] if futuros else max(jogos, key=lambda jogo: (jogo.data, jogo.id))
        return adversario, jogo_referencia.equipa, adversario.nome
    if tipo == "plano-jogo":
        jogo = db.query(Jogo).filter(Jogo.id == alvo_id, Jogo.clube_id == clube_id).first()
        if not jogo:
            raise HTTPException(404, "Jogo nao encontrado.")
        verificar_ou_negar_equipa(utilizador, jogo.equipa_id, db)
        if not jogo_em_preparacao(jogo):
            raise HTTPException(409, "Nao e possivel exportar um plano de jogo ja realizado.")
        return jogo, jogo.equipa, f"{jogo.equipa.nome} vs {jogo.adversario}"
    raise HTTPException(404, "Tipo de relatorio desconhecido.")


def _itens_por_categoria(itens) -> dict[str, list]:
    resultado: dict[str, list] = {}
    for relacao in itens or []:
        categoria = getattr(getattr(relacao, "item", None), "categoria", "outros")
        resultado.setdefault(categoria, []).append(relacao)
    return resultado


def _jogadores_prioritarios_coletivo(adversario, jogador_ids: list[int] | None = None) -> list[dict]:
    """Resume os jogadores escolhidos para os mini-cartoes do relatorio."""
    resultado = []
    vistos = set()
    if jogador_ids is None:
        candidatos = [
            (lider.jogador, lider.tipo)
            for lider in (adversario.lideres or [])
            if lider.jogador
        ]
    else:
        jogadores = {jogador.id: jogador for jogador in (adversario.jogadores or [])}
        candidatos = [(jogadores.get(jogador_id), None) for jogador_id in jogador_ids]

    for jogador, caracteristica_padrao in candidatos:
        if not jogador or jogador.id in vistos:
            continue
        vistos.add(jogador.id)
        caracteristicas = [
            relacao for relacao in (jogador.itens or [])
            if relacao.item
            and relacao.item.ativo
            and relacao.item.categoria == "caracteristica"
            and relacao.avaliacao is not None
        ]
        caracteristicas.sort(key=lambda relacao: (
            -(relacao.avaliacao or 0),
            not relacao.destaque,
            relacao.ordem_destaque if relacao.ordem_destaque is not None else 10_000,
            relacao.item.nome.casefold(),
        ))
        resultado.append({
            "jogador": jogador,
            "caracteristica": (
                caracteristicas[0].item.nome
                if caracteristicas
                else (caracteristica_padrao or "Sem caracteristica avaliada")
            ),
        })
    return resultado[:4]


def _caracteristicas_cinco_coletivo(jogadores) -> dict[int, list[str]]:
    resultado = {}
    for jogador in jogadores:
        itens = [
            relacao for relacao in (jogador.itens or [])
            if relacao.item and relacao.item.ativo
            and relacao.item.categoria == "caracteristica"
            and (relacao.avaliacao or 0) > 0
        ]
        itens.sort(key=lambda relacao: (
            not relacao.destaque,
            -(relacao.avaliacao or 0),
            relacao.ordem_destaque if relacao.ordem_destaque is not None else 10_000,
        ))
        resultado[jogador.id] = [relacao.item.nome for relacao in itens[:2]]
    return resultado


def _opcoes_defensivas_coletivas(opcoes) -> list[dict[str, str]]:
    if not opcoes:
        return []
    campos = (
        ("bloqueio_central", "Bloqueios diretos centrais"),
        ("bloqueio_lateral", "Bloqueios diretos laterais"),
        ("bloqueio_indireto", "Bloqueios indiretos"),
        ("handoff", "Handoff"),
        ("bola_interior", "Defesa da bola interior"),
        ("alteracoes", "Alterações defensivas"),
    )
    return [
        {"nome": nome, "notas": (getattr(opcoes, f"{campo}_notas", None) or "").strip()}
        for campo, nome in campos
        if getattr(opcoes, campo, False) and (getattr(opcoes, f"{campo}_notas", None) or "").strip()
    ]


def _rotacoes_coletivas_por_posicao(adversario):
    cinco_previsto = [r for r in adversario.rotacao_prevista if r.papel == "titular"]
    rotacoes_previstas = [r for r in adversario.rotacao_prevista if r.papel != "titular"]
    rotacoes_por_posicao = {}
    for entrada in cinco_previsto + rotacoes_previstas:
        posicao = (entrada.jogador.posicao_principal or "Sem posição").strip()
        rotacoes_por_posicao.setdefault(posicao, [])
    for entrada in rotacoes_previstas:
        posicao = (entrada.jogador.posicao_principal or "Sem posição").strip()
        rotacoes_por_posicao[posicao].append(entrada)
    return cinco_previsto, rotacoes_previstas, rotacoes_por_posicao


def _substituicoes_coletivas_por_titular(adversario, titulares_ids=None):
    resultado = {}
    for entrada in adversario.substituicoes_previstas:
        if titulares_ids is not None and entrada.titular_id not in titulares_ids:
            continue
        resultado.setdefault(entrada.titular_id, {})[entrada.ordem] = entrada
    return resultado


def _media_exibida_coletiva(adversario) -> dict[str, float | None]:
    media = adversario.media_coletiva
    if media and media.usar_manual:
        return {campo: getattr(media, campo) for campo in STAT_FIELDS}
    jogos = adversario.estatisticas_coletivas or []
    resultado = {}
    for campo in STAT_FIELDS:
        valores = [getattr(jogo, campo) for jogo in jogos if getattr(jogo, campo) is not None]
        resultado[campo] = round(sum(valores) / len(valores), 1) if valores else None
    return resultado


def _indicadores_coletivos(media: dict, selecionados: list[str]) -> list[dict]:
    padrao = ["pontos", "dois_c", "dois_t", "tres_c", "tres_t", "ll_c", "ll_t", "ress_of", "ress_def", "assistencias", "turnovers", "roubos", "eficiencia"]
    campos = selecionados or padrao
    indicadores = []
    pares = {"fg_c": ("fg_t", "FG%"), "dois_c": ("dois_t", "2P%"), "tres_c": ("tres_t", "3P%"), "ll_c": ("ll_t", "LL%")}
    usados = set()
    for campo in campos:
        if campo in usados or media.get(campo) is None:
            continue
        if campo in pares:
            tentativas, nome = pares[campo]
            if tentativas in campos and media.get(tentativas):
                indicadores.append({"nome": nome, "valor": f"{media[campo] / media[tentativas] * 100:.0f}%"})
                usados.update((campo, tentativas))
                continue
        if any(campo == tentativas and convertidos in campos and media.get(convertidos) is not None for convertidos, (tentativas, _) in pares.items()):
            continue
        indicadores.append({"nome": STAT_CODES[campo], "valor": f"{media[campo]:.1f}"})
        usados.add(campo)
    return indicadores


def _destaque_estatistico_coletivo(codigo, tom, indicadores: list[dict]) -> dict | None:
    codigo = str(codigo or "").strip()
    tom = str(tom or "").strip().lower()
    if codigo in STAT_HIGHLIGHT_CODES and tom in {"positivo", "negativo"}:
        if any(indicador["nome"] == codigo for indicador in indicadores):
            return {"codigo": codigo, "tom": tom}
    return None


def _acoes_principais_coletivas(adversario) -> list:
    return [entrada for entrada in adversario.conceitos_ofensivos if entrada.setplay]


def _grupos_resposta_coletiva(blocos: dict[str, list]) -> list[dict]:
    grupos = [
        ("Sem bola", "resposta", ("resposta",)),
        ("Prioridades", "prioridades", ("prioridade_ofensiva", "prioridade_defensiva")),
        ("Transição", "situacoes", ("transicao",)),
        ("Ressalto", "situacoes", ("ressalto",)),
        ("Situações especiais", "situacoes", ("situacao_especial",)),
    ]
    resultado = []
    for nome, secao, tipos in grupos:
        candidatos = [bloco for tipo in tipos for bloco in blocos.get(tipo, []) if bloco.destaque]
        resultado.append({"nome": nome, "secao": secao, "blocos": candidatos})
    return resultado


def _contexto_relatorio(tipo: str, alvo, equipa: Equipa, dados, db: Session) -> dict:
    contexto = {"tipo_relatorio": tipo, "equipa": equipa, "alvo": alvo}
    if tipo in {"roster", "scouting-individual", "scouting-coletivo"}:
        contexto["simbolo_adversario"] = _fotografia_embebida(alvo.simbolo)
    ids = {int(v) for v in dados.getlist("entidades") if str(v).isdigit()}
    stat_fields = [campo for campo in dados.getlist("estatisticas") if campo in STAT_FIELDS]
    if tipo == "roster":
        atletas = [j for j in alvo.jogadores if j.estado != "arquivado" and (not ids or j.id in ids)]
        contexto.update(
            atletas=sorted(atletas, key=lambda j: j.nome.lower()),
            fotografias={j.id: _fotografia_embebida(j.fotografia) for j in atletas},
            stat_fields=stat_fields,
            stat_labels=STAT_LABELS,
            stat_codes=STAT_CODES,
        )
    elif tipo == "scouting-individual":
        atletas = [j for j in alvo.jogadores if not ids or j.id in ids]
        contexto.update(
            atletas=atletas,
            fotografias={j.id: _fotografia_embebida(j.fotografia) for j in atletas},
            itens_atletas={j.id: _itens_por_categoria(j.itens) for j in atletas},
            perfis_atletas={j.id: construir_perfil_individual(j) for j in atletas},
            destaques_atletas={j.id: construir_perfil_individual(j, apenas_destaques=True) for j in atletas},
            ultimo_jogo={j.id: (j.estatisticas[0] if j.estatisticas else None) for j in atletas},
            campo_meio=imagem_estatica_data_uri(STATIC_DIR / "campos" / "campo_meio.png"),
            stat_fields=stat_fields,
            stat_labels=STAT_LABELS,
            stat_codes=STAT_CODES,
        )
    elif tipo == "scouting-coletivo":
        jogadores_destaque_ids = None
        if dados.get("configuracao_jogadores_destaque"):
            jogadores_destaque_ids = []
            for valor in dados.getlist("jogadores_destaque"):
                if str(valor).isdigit() and int(valor) not in jogadores_destaque_ids:
                    jogadores_destaque_ids.append(int(valor))
                if len(jogadores_destaque_ids) == 4:
                    break
        setplays = [r.setplay for r in alvo.conceitos_ofensivos]
        itens_coletivos = _itens_por_categoria(alvo.itens_coletivos)
        media_exibida = _media_exibida_coletiva(alvo)
        blocos_coletivos = {}
        for bloco in alvo.blocos_coletivos:
            blocos_coletivos.setdefault(bloco.tipo, []).append(bloco)
        perfil_preenchido = [
            bloco for bloco in blocos_coletivos.get("perfil", [])
            if (bloco.avaliacao > 0 if bloco.avaliacao is not None else bool(bloco.nivel))
        ]
        perfil_facil = []
        for bloco in perfil_preenchido:
            avaliacao = bloco.avaliacao if bloco.avaliacao is not None else {"fraco": 2, "medio": 3, "forte": 4}.get(bloco.nivel)
            if bloco.destaque and avaliacao not in (None, 0):
                perfil_facil.append({
                    "titulo": bloco.titulo, "descricao": bloco.descricao,
                    "categoria": bloco.categoria, "avaliacao": avaliacao,
                })
        jogos_referencia = jogos_em_preparacao(
            [ligacao.jogo for ligacao in alvo.jogos if ligacao.jogo]
        )
        indicadores_coletivos = _indicadores_coletivos(media_exibida, stat_fields)
        destaque_estatistico = _destaque_estatistico_coletivo(
            dados.get("estatistica_destaque"), dados.get("estatistica_destaque_tom"), indicadores_coletivos
        )
        cinco_previsto, rotacoes_previstas, rotacoes_por_posicao = _rotacoes_coletivas_por_posicao(alvo)
        substituicoes_por_titular = _substituicoes_coletivas_por_titular(
            alvo, {entrada.jogador_id for entrada in cinco_previsto}
        )
        ids_substitutos = {
            entrada.jogador_id
            for entradas in substituicoes_por_titular.values()
            for entrada in entradas.values()
        }
        rotacoes_sem_titular = {
            posicao: [entrada for entrada in entradas if entrada.jogador_id not in ids_substitutos]
            for posicao, entradas in rotacoes_por_posicao.items()
        }
        contexto.update(
            adversario=alvo,
            jogo_referencia=jogos_referencia[0] if jogos_referencia else None,
            fotografias={j.id: _fotografia_embebida(j.fotografia) for j in alvo.jogadores},
            caracteristicas_cinco=_caracteristicas_cinco_coletivo(
                [entrada.jogador for entrada in cinco_previsto]
            ),
            opcoes_defensivas_facil=_opcoes_defensivas_coletivas(alvo.opcoes_defensivas),
            itens_coletivos=itens_coletivos,
            jogadas_coletivas=_acoes_principais_coletivas(alvo),
            jogadores_prioritarios=_jogadores_prioritarios_coletivo(alvo, jogadores_destaque_ids),
            media_exibida=media_exibida,
            indicadores_coletivos=indicadores_coletivos,
            destaque_estatistico=destaque_estatistico,
            perfil_ataque=[bloco for bloco in perfil_facil if "defesa" not in (bloco["categoria"] or "").lower()],
            perfil_defesa=[bloco for bloco in perfil_facil if "defesa" in (bloco["categoria"] or "").lower()],
            respostas_grupos=_grupos_resposta_coletiva(blocos_coletivos),
            stat_fields_com_dados=[
                campo for campo in stat_fields
                if media_exibida.get(campo) is not None
            ],
            caracteristicas_estatisticas_com_dados=[
                relacao for relacao in itens_coletivos.get("estatistica_coletiva", [])
                if str(relacao.valor or "").strip() or str(relacao.observacao or "").strip()
            ],
            principais_caracteristicas=(
                itens_coletivos.get("ataque_forte", [])
                + itens_coletivos.get("defesa_forte", [])
            ),
            blocos_coletivos=blocos_coletivos,
            perfil_preenchido=perfil_preenchido,
            cinco_previsto=cinco_previsto,
            rotacoes_previstas=rotacoes_previstas,
            rotacoes_por_posicao=rotacoes_por_posicao,
            rotacoes_sem_titular=rotacoes_sem_titular,
            rotacoes_sem_titular_lista=[entrada for entradas in rotacoes_sem_titular.values() for entrada in entradas],
            ha_rotacoes_sem_titular=any(rotacoes_sem_titular.values()),
            substituicoes_por_titular=substituicoes_por_titular,
            passos_setplay={setplay.id: _passos_setplay_pdf(setplay) for setplay in setplays},
            stat_fields=stat_fields,
            stat_labels=STAT_LABELS,
            stat_codes=STAT_CODES,
        )
    else:
        plano = alvo.plano
        titulares_ids = [getattr(plano, f"titular_{i}", None) for i in range(1, 6)] if plano else []
        titulares = db.query(Jogador).filter(Jogador.id.in_([i for i in titulares_ids if i])).all() if any(titulares_ids) else []
        titulares_map = {j.id: j for j in titulares}
        titulares_ordenados = [titulares_map[i] for i in titulares_ids if i in titulares_map]
        ligacao = db.query(ScoutingAdversarioJogo).filter(ScoutingAdversarioJogo.jogo_id == alvo.id).first()
        adversario = ligacao.adversario if ligacao else None
        contexto.update(
            jogo=alvo,
            plano=plano,
            convocados=alvo.convocados,
            titulares=titulares_ordenados,
            rotacoes_equipa=rotacoes_por_grupo(plano),
            graficos_plano=graficos_plano(plano),
            fotografias_equipa={
                j.id: _fotografia_embebida(j.fotografia)
                for j in {c.jogador for c in alvo.convocados} | set(titulares_ordenados)
            },
            adversario=adversario,
            simbolo_adversario=_fotografia_embebida(adversario.simbolo) if adversario else None,
            fotografias_adversario={
                j.id: _fotografia_embebida(j.fotografia)
                for j in (adversario.jogadores if adversario else [])
            },
            cinco_adversario=[r for r in adversario.rotacao_prevista if r.papel == "titular"] if adversario else [],
            itens_coletivos=_itens_por_categoria(adversario.itens_coletivos) if adversario else {},
        )
    return contexto


def _predefinicoes_relatorio(db: Session, utilizador, tipo: str, equipa_id: int) -> list[PredefinicaoRelatorio]:
    return db.query(PredefinicaoRelatorio).filter(
        PredefinicaoRelatorio.clube_id == utilizador.clube_id,
        PredefinicaoRelatorio.tipo == tipo,
        or_(
            and_(PredefinicaoRelatorio.ambito == "pessoal", PredefinicaoRelatorio.utilizador_id == utilizador.id),
            and_(PredefinicaoRelatorio.ambito == "equipa", PredefinicaoRelatorio.equipa_id == equipa_id),
        ),
    ).order_by(PredefinicaoRelatorio.ambito, PredefinicaoRelatorio.nome).all()


def _contexto_configurador(tipo: str, alvo_id: int, db: Session, utilizador, predefinicao_id: int | None = None) -> dict:
    alvo, equipa, alvo_nome = _alvo_relatorio(tipo, alvo_id, db, utilizador)
    predefinicoes = _predefinicoes_relatorio(db, utilizador, tipo, equipa.id)
    predefinicao = next((p for p in predefinicoes if p.id == predefinicao_id), None)
    formato_inicial = "facil" if tipo == "scouting-coletivo" else "completo"
    selecionadas = json.loads(predefinicao.secoes) if predefinicao else _secoes_padrao(tipo, formato_inicial)
    if tipo == "scouting-coletivo":
        selecionadas = _normalizar_secoes_coletivas(selecionadas)
    try:
        configuracao_guardada = json.loads(predefinicao.configuracao or "{}") if predefinicao else {}
    except (TypeError, ValueError):
        configuracao_guardada = {}
    entidades = []
    if tipo == "roster":
        entidades = [j for j in alvo.jogadores if j.estado != "arquivado"]
    elif tipo == "scouting-individual":
        entidades = alvo.jogadores
    jogadores_mini_cartoes = []
    jogadores_mini_cartoes_preselecionados = []
    if tipo == "scouting-coletivo":
        jogadores_mini_cartoes = sorted(alvo.jogadores, key=lambda jogador: jogador.nome.casefold())
        jogadores_mini_cartoes_preselecionados = [
            entrada["jogador"].id for entrada in _jogadores_prioritarios_coletivo(alvo)
        ]
    return {
        "utilizador": utilizador,
        "tipo": tipo,
        "alvo_id": alvo_id,
        "alvo_nome": alvo_nome,
        "equipa": equipa,
        "config": RELATORIOS_CONFIG[tipo],
        "formatos": FORMATOS_RELATORIO,
        "entidades": entidades,
        "jogadores_mini_cartoes": jogadores_mini_cartoes,
        "jogadores_mini_cartoes_preselecionados": jogadores_mini_cartoes_preselecionados,
        "predefinicoes": predefinicoes,
        "predefinicao": predefinicao,
        "selecionadas": selecionadas,
        "estatisticas_selecionadas": configuracao_guardada.get("estatisticas", []),
        "estatisticas_destaque_opcoes": STAT_HIGHLIGHT_OPTIONS if tipo == "scouting-coletivo" else [],
        "estatistica_destaque_selecionada": configuracao_guardada.get("estatistica_destaque", "") if configuracao_guardada.get("estatistica_destaque") in STAT_HIGHLIGHT_CODES else "",
        "estatistica_destaque_tom": configuracao_guardada.get("estatistica_destaque_tom", "positivo") if configuracao_guardada.get("estatistica_destaque_tom") in {"positivo", "negativo"} else "positivo",
        "jogadores_destaque_selecionados": configuracao_guardada.get("jogadores_destaque", []),
        "padroes": {f: _secoes_padrao(tipo, f) for f in FORMATOS_RELATORIO},
        "estatisticas_disponiveis": [(campo, STAT_LABELS[campo]) for campo in STAT_FIELDS]
        if any(secao["chave"] == "estatisticas" for secao in RELATORIOS_CONFIG[tipo]["secoes"])
        else [],
        "pode_ver_contactos": _perfil(utilizador) in PERFIS_CONTACTOS,
    }


@router.get("/relatorios/{tipo}/{alvo_id}", response_class=HTMLResponse)
def configurar_relatorio(
    tipo: str,
    alvo_id: int,
    request: Request,
    predefinicao_id: int | None = None,
    db: Session = Depends(get_db),
    utilizador=Depends(exigir_operacao_desportiva),
):
    if tipo not in RELATORIOS_CONFIG:
        raise HTTPException(404, "Tipo de relatorio desconhecido.")
    return templates.TemplateResponse(
        request,
        "exportacoes/relatorio_configurar.html",
        _contexto_configurador(tipo, alvo_id, db, utilizador, predefinicao_id),
    )


@router.get("/relatorios/{tipo}/{alvo_id}/fragmento", response_class=HTMLResponse)
def fragmento_configurar_relatorio(
    tipo: str,
    alvo_id: int,
    request: Request,
    predefinicao_id: int | None = None,
    db: Session = Depends(get_db),
    utilizador=Depends(exigir_operacao_desportiva),
):
    if tipo not in RELATORIOS_CONFIG:
        raise HTTPException(404, "Tipo de relatorio desconhecido.")
    return templates.TemplateResponse(
        request,
        "exportacoes/_relatorio_formulario.html",
        _contexto_configurador(tipo, alvo_id, db, utilizador, predefinicao_id),
    )


def _gerar_relatorio(tipo: str, alvo_id: int, request: Request, dados, db: Session, utilizador, previsualizar: bool):
    if tipo not in RELATORIOS_CONFIG:
        raise HTTPException(404, "Tipo de relatorio desconhecido.")
    alvo, equipa, alvo_nome = _alvo_relatorio(tipo, alvo_id, db, utilizador)
    formato, orientacao, secoes = _configuracao_relatorio(tipo, dados, utilizador)
    contexto = _contexto_relatorio(tipo, alvo, equipa, dados, db)
    if tipo == "scouting-individual":
        secoes_por_atleta, estatisticas_por_atleta = _configuracao_individual_por_atleta(
            dados, contexto["atletas"], secoes
        )
        contexto.update(
            secoes_por_atleta=secoes_por_atleta,
            stat_fields_por_atleta=estatisticas_por_atleta,
        )
    if tipo == "scouting-coletivo":
        contexto["respostas_grupos"] = [
            grupo for grupo in contexto["respostas_grupos"] if grupo["secao"] in secoes
        ]
    contexto.update(
        formato=formato,
        orientacao=orientacao,
        secoes=secoes,
        secoes_set=set(secoes),
        titulo_relatorio=RELATORIOS_CONFIG[tipo]["nome"],
        alvo_nome=alvo_nome,
        nome_autor=utilizador.nome,
        pode_ver_contactos=_perfil(utilizador) in PERFIS_CONTACTOS and formato != "a3",
        ordem_secoes={chave: indice for indice, chave in enumerate(secoes)},
    )
    template = f"relatorios/{tipo.replace('-', '_')}_{formato}.html"
    if tipo == "scouting-coletivo" and formato == "facil":
        template = "relatorios/scouting_coletivo_facil_v2.html"
    nome = f"{RELATORIOS_CONFIG[tipo]['nome']} - {alvo_nome} - {date.today().isoformat()}"
    if previsualizar:
        html = templates.get_template(template).render(preparar_contexto_pdf(request, contexto))
        return HTMLResponse(html)
    pdf = renderizar_pdf(request, template, contexto)
    return Response(pdf, media_type="application/pdf", headers={"Content-Disposition": content_disposition_pdf(nome)})


@router.post("/relatorios/{tipo}/{alvo_id}/previsualizar")
async def previsualizar_relatorio(tipo: str, alvo_id: int, request: Request, db: Session = Depends(get_db), utilizador=Depends(exigir_operacao_desportiva)):
    return _gerar_relatorio(tipo, alvo_id, request, await request.form(), db, utilizador, True)


@router.post("/relatorios/{tipo}/{alvo_id}/pdf")
async def descarregar_relatorio(tipo: str, alvo_id: int, request: Request, db: Session = Depends(get_db), utilizador=Depends(exigir_operacao_desportiva)):
    return _gerar_relatorio(tipo, alvo_id, request, await request.form(), db, utilizador, False)


@router.post("/relatorios/{tipo}/{alvo_id}/predefinicoes")
async def guardar_predefinicao_relatorio(tipo: str, alvo_id: int, request: Request, db: Session = Depends(get_db), utilizador=Depends(exigir_operacao_desportiva)):
    if tipo not in RELATORIOS_CONFIG:
        raise HTTPException(404, "Tipo de relatorio desconhecido.")
    _, equipa, _ = _alvo_relatorio(tipo, alvo_id, db, utilizador)
    dados = await request.form()
    nome = str(dados.get("nome_predefinicao", "")).strip()
    ambito = str(dados.get("ambito", "pessoal"))
    if not nome or ambito not in {"pessoal", "equipa"}:
        raise HTTPException(400, "Indica um nome e um ambito validos.")
    formato, orientacao, secoes = _configuracao_relatorio(tipo, dados, utilizador)
    filtro = [
        PredefinicaoRelatorio.clube_id == utilizador.clube_id,
        PredefinicaoRelatorio.tipo == tipo,
        PredefinicaoRelatorio.ambito == ambito,
        PredefinicaoRelatorio.nome == nome,
    ]
    if ambito == "pessoal":
        filtro.append(PredefinicaoRelatorio.utilizador_id == utilizador.id)
    else:
        filtro.append(PredefinicaoRelatorio.equipa_id == equipa.id)
    registo = db.query(PredefinicaoRelatorio).filter(*filtro).first()
    if not registo:
        registo = PredefinicaoRelatorio(
            clube_id=utilizador.clube_id,
            utilizador_id=utilizador.id if ambito == "pessoal" else None,
            equipa_id=equipa.id if ambito == "equipa" else None,
            tipo=tipo,
            ambito=ambito,
            nome=nome,
            protegido=False,
        )
        db.add(registo)
    registo.formato = formato
    registo.orientacao = orientacao
    registo.secoes = json.dumps(secoes)
    registo.configuracao = json.dumps({
        "estatisticas": [campo for campo in dados.getlist("estatisticas") if campo in STAT_FIELDS],
        "estatistica_destaque": str(dados.get("estatistica_destaque") or "").strip() if tipo == "scouting-coletivo" and str(dados.get("estatistica_destaque") or "").strip() in STAT_HIGHLIGHT_CODES else "",
        "estatistica_destaque_tom": str(dados.get("estatistica_destaque_tom") or "").strip().lower() if tipo == "scouting-coletivo" and str(dados.get("estatistica_destaque_tom") or "").strip().lower() in {"positivo", "negativo"} else "positivo",
        "jogadores_destaque": [int(valor) for valor in dados.getlist("jogadores_destaque") if str(valor).isdigit()][:4],
    })
    db.commit()
    return RedirectResponse(f"/exportacoes/relatorios/{tipo}/{alvo_id}?predefinicao_id={registo.id}", status_code=303)
