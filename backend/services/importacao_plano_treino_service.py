"""Leitura e interpretação conservadora de planos de treino enviados pelo treinador."""

from __future__ import annotations

from datetime import date
from io import BytesIO
import re
import unicodedata
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageEnhance, ImageFilter, ImageOps
from pypdf import PdfReader


MAX_FICHEIRO_BYTES = 12 * 1024 * 1024
MAX_PAGINAS = 12
MAX_TEXTO_CARACTERES = 500_000
MAX_CONTEUDO_DESCOMPRIMIDO = 80 * 1024 * 1024
CATEGORIAS_EXERCICIO = ("TIO", "TID", "TaIO", "TaID", "TCO", "TCD", "FIS", "PSI")
EXTENSOES_ACEITES = {".pdf", ".png", ".jpg", ".jpeg", ".docx", ".xlsx"}
TIPOS_MIME_SEGUROS = {
    ".pdf": "application/pdf",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}

CAMPOS_CATEGORIA = {
    "TIO": "min_tecnica_individual_ofensiva",
    "TID": "min_tecnica_individual_defensiva",
    "TaIO": "min_tatica_individual_ofensiva",
    "TaID": "min_tatica_individual_defensiva",
    "TCO": "min_tatica_coletiva_ofensiva",
    "TCD": "min_tatica_coletiva_defensiva",
    "FIS": "min_fisicos",
    "PSI": "min_psicologicos",
}

_ALIAS_CATEGORIA = {
    "TIO": "TIO",
    "TECNICA INDIVIDUAL OFENSIVA": "TIO",
    "TID": "TID",
    "TECNICA INDIVIDUAL DEFENSIVA": "TID",
    "TAIO": "TaIO",
    "TATICA INDIVIDUAL OFENSIVA": "TaIO",
    "TAID": "TaID",
    "TATICA INDIVIDUAL DEFENSIVA": "TaID",
    "TCO": "TCO",
    "TATICA COLETIVA OFENSIVA": "TCO",
    "TCD": "TCD",
    "TATICA COLETIVA DEFENSIVA": "TCD",
    "FIS": "FIS",
    "FISICO": "FIS",
    "FISICA": "FIS",
    "PREPARACAO FISICA": "FIS",
    "PSI": "PSI",
    "PSICOLOGIA": "PSI",
    "PSICOLOGICO": "PSI",
}


class ErroImportacaoPlano(ValueError):
    pass


def _validar_documento_comprimido(conteudo: bytes) -> None:
    try:
        with ZipFile(BytesIO(conteudo)) as arquivo:
            entradas = arquivo.infolist()
            if len(entradas) > 2500 or sum(item.file_size for item in entradas) > MAX_CONTEUDO_DESCOMPRIMIDO:
                raise ErroImportacaoPlano("O documento comprimido é demasiado complexo ou volumoso.")
            if any(item.file_size > 25 * 1024 * 1024 for item in entradas):
                raise ErroImportacaoPlano("O documento contém um elemento interno demasiado volumoso.")
    except BadZipFile as erro:
        raise ErroImportacaoPlano("O documento Word ou Excel está danificado.") from erro


def validar_ficheiro(nome: str, tipo_mime: str | None, conteudo: bytes) -> tuple[str, str]:
    nome_seguro = Path(nome or "").name.strip()
    extensao = Path(nome_seguro).suffix.lower()
    if not nome_seguro or extensao not in EXTENSOES_ACEITES:
        raise ErroImportacaoPlano("Usa um ficheiro PDF, Word, Excel, PNG ou JPEG.")
    if not conteudo:
        raise ErroImportacaoPlano("O ficheiro está vazio.")
    if len(conteudo) > MAX_FICHEIRO_BYTES:
        raise ErroImportacaoPlano("O ficheiro não pode ultrapassar 12 MB.")
    if extensao == ".pdf" and not conteudo.startswith(b"%PDF"):
        raise ErroImportacaoPlano("O conteúdo não corresponde a um PDF válido.")
    if extensao == ".png" and not conteudo.startswith(b"\x89PNG\r\n\x1a\n"):
        raise ErroImportacaoPlano("O conteúdo não corresponde a uma imagem PNG válida.")
    if extensao in {".jpg", ".jpeg"} and not conteudo.startswith(b"\xff\xd8"):
        raise ErroImportacaoPlano("O conteúdo não corresponde a uma imagem JPEG válida.")
    if extensao in {".docx", ".xlsx"} and not conteudo.startswith(b"PK"):
        raise ErroImportacaoPlano("O conteúdo não corresponde ao documento indicado.")
    if extensao in {".docx", ".xlsx"}:
        _validar_documento_comprimido(conteudo)
    return nome_seguro, TIPOS_MIME_SEGUROS[extensao]


def _texto_documento_word(conteudo: bytes) -> str:
    documento = Document(BytesIO(conteudo))
    linhas = [paragrafo.text.strip() for paragrafo in documento.paragraphs if paragrafo.text.strip()]
    for tabela in documento.tables:
        for linha in tabela.rows:
            celulas = [celula.text.strip().replace("\n", " ") for celula in linha.cells]
            if any(celulas):
                linhas.append(" | ".join(celulas))
    return "\n".join(linhas)


def _texto_excel(conteudo: bytes) -> str:
    livro = load_workbook(BytesIO(conteudo), read_only=True, data_only=True)
    linhas = []
    for folha in livro.worksheets[:50]:
        linhas.append(f"Folha: {folha.title}")
        for indice, valores in enumerate(folha.iter_rows(values_only=True), start=1):
            if indice > 5000:
                linhas.append("[Restantes linhas omitidas por segurança]")
                break
            celulas = [str(valor).strip() if valor is not None else "" for valor in valores[:200]]
            if any(celulas):
                linhas.append(" | ".join(celulas))
    livro.close()
    return "\n".join(linhas)


def _preparar_imagem(imagem: Image.Image) -> Image.Image:
    if imagem.width * imagem.height > 30_000_000:
        raise ErroImportacaoPlano("A imagem não pode ultrapassar 30 milhões de píxeis.")
    imagem = ImageOps.exif_transpose(imagem)
    if max(imagem.size) > 3000:
        imagem.thumbnail((3000, 3000), Image.Resampling.LANCZOS)
    cinzento = ImageOps.grayscale(imagem)
    cinzento = ImageOps.autocontrast(cinzento, cutoff=1)
    cinzento = ImageEnhance.Contrast(cinzento).enhance(1.25)
    return cinzento.filter(ImageFilter.SHARPEN)


def _ocr_imagem(imagem: Image.Image) -> tuple[str, float | None]:
    try:
        import pytesseract
        from pytesseract import Output
    except ImportError as erro:
        raise ErroImportacaoPlano("O reconhecimento de imagens não está instalado no servidor.") from erro

    preparada = _preparar_imagem(imagem)
    ultima_excecao = None
    for idioma in ("por+eng", "eng"):
        try:
            dados = pytesseract.image_to_data(
                preparada,
                lang=idioma,
                config="--oem 1 --psm 6",
                output_type=Output.DICT,
            )
            palavras = []
            confiancas = []
            linhas: dict[tuple[int, int, int], list[str]] = {}
            for indice, texto in enumerate(dados.get("text", [])):
                texto = (texto or "").strip()
                if not texto:
                    continue
                chave = (
                    int(dados["block_num"][indice]),
                    int(dados["par_num"][indice]),
                    int(dados["line_num"][indice]),
                )
                linhas.setdefault(chave, []).append(texto)
                palavras.append(texto)
                try:
                    valor = float(dados["conf"][indice])
                    if valor >= 0:
                        confiancas.append(valor)
                except (TypeError, ValueError):
                    pass
            texto_final = "\n".join(" ".join(linhas[chave]) for chave in sorted(linhas))
            confianca = sum(confiancas) / len(confiancas) if confiancas else None
            return texto_final.strip(), confianca
        except Exception as erro:  # Tesseract ausente ou pacote de idioma indisponível.
            ultima_excecao = erro
    raise ErroImportacaoPlano(
        "Não foi possível ler a imagem. Confirma que está direita, focada e com boa luz."
    ) from ultima_excecao


def _texto_pdf(conteudo: bytes) -> tuple[str, float | None, list[str]]:
    leitor = PdfReader(BytesIO(conteudo))
    if len(leitor.pages) > MAX_PAGINAS:
        raise ErroImportacaoPlano(f"O PDF não pode ter mais de {MAX_PAGINAS} páginas.")
    textos = [(pagina.extract_text() or "").strip() for pagina in leitor.pages]
    avisos = []
    confiancas = []
    paginas_ocr = [indice for indice, texto in enumerate(textos) if len(texto) < 35]
    if paginas_ocr:
        try:
            import pypdfium2 as pdfium

            documento = pdfium.PdfDocument(conteudo)
            for indice in paginas_ocr:
                pagina = documento[indice]
                bitmap = pagina.render(scale=2.4)
                texto_ocr, confianca = _ocr_imagem(bitmap.to_pil())
                textos[indice] = texto_ocr
                if confianca is not None:
                    confiancas.append(confianca)
        except ErroImportacaoPlano as erro:
            avisos.append(str(erro))
        except Exception:
            avisos.append("Algumas páginas digitalizadas não puderam ser reconhecidas automaticamente.")
    confianca_media = sum(confiancas) / len(confiancas) if confiancas else (100.0 if any(textos) else None)
    return "\n\n".join(texto for texto in textos if texto), confianca_media, avisos


def extrair_texto(nome: str, conteudo: bytes) -> tuple[str, float | None, list[str]]:
    extensao = Path(nome).suffix.lower()
    try:
        if extensao == ".pdf":
            texto, confianca, avisos = _texto_pdf(conteudo)
            return texto[:MAX_TEXTO_CARACTERES], confianca, avisos
        if extensao == ".docx":
            return _texto_documento_word(conteudo)[:MAX_TEXTO_CARACTERES], 100.0, []
        if extensao == ".xlsx":
            return _texto_excel(conteudo)[:MAX_TEXTO_CARACTERES], 100.0, []
        if extensao in {".png", ".jpg", ".jpeg"}:
            imagem = Image.open(BytesIO(conteudo))
            imagem.verify()
            imagem = Image.open(BytesIO(conteudo))
            texto, confianca = _ocr_imagem(imagem)
            return texto[:MAX_TEXTO_CARACTERES], confianca, []
    except ErroImportacaoPlano:
        raise
    except Exception as erro:
        raise ErroImportacaoPlano("Não foi possível abrir o documento enviado.") from erro
    raise ErroImportacaoPlano("Formato de ficheiro não suportado.")


def _sem_acentos(valor: str) -> str:
    return "".join(
        caracter for caracter in unicodedata.normalize("NFD", valor.upper())
        if unicodedata.category(caracter) != "Mn"
    )


def normalizar_categoria(valor: str | None) -> str | None:
    normalizado = _sem_acentos((valor or "").strip())
    if not normalizado:
        return None
    if normalizado in _ALIAS_CATEGORIA:
        return _ALIAS_CATEGORIA[normalizado]
    for alias in sorted(_ALIAS_CATEGORIA, key=len, reverse=True):
        if re.search(rf"(?<!\w){re.escape(alias)}(?!\w)", normalizado):
            return _ALIAS_CATEGORIA[alias]
    return None


def _primeiro_grupo(padrao: str, texto: str, flags: int = re.IGNORECASE) -> str:
    correspondencia = re.search(padrao, texto, flags)
    return correspondencia.group(1).strip() if correspondencia else ""


def _extrair_minutos(linha: str, categoria: str | None) -> int | None:
    padroes = [
        r"(?<!\d)(\d{1,3})\s*(?:min(?:uto)?s?|mn|')(?!\w)",
        r"(?:dura[cç][aã]o|tempo)\s*[:\-]?\s*(\d{1,3})(?!\d)",
    ]
    if categoria:
        padroes.extend([
            rf"{re.escape(categoria)}\s*[|:\-]?\s*(\d{{1,3}})(?![:\d])",
            rf"(?<![:\d])(\d{{1,3}})\s*[|:\-]?\s*{re.escape(categoria)}",
        ])
    for padrao in padroes:
        valor = _primeiro_grupo(padrao, linha)
        if valor and 1 <= int(valor) <= 240:
            return int(valor)
    if "|" in linha and categoria:
        numeros = [int(item) for item in re.findall(r"(?<![:\d])(\d{1,3})(?![:\d])", linha)]
        candidatos = [numero for numero in numeros if 1 <= numero <= 240]
        if candidatos:
            return candidatos[0]
    return None


def _limpar_nome_exercicio(linha: str, categoria: str | None, minutos: int | None) -> str:
    nome = linha
    nome = re.sub(r"\b(?:exerc[ií]cio|ex\.?|drill)\s*\d*\s*[:\-]?", "", nome, flags=re.IGNORECASE)
    nome = re.sub(r"\bRPE\s*[:\-]?\s*(?:10|[0-9])(?:[,.]5)?\b", "", nome, flags=re.IGNORECASE)
    nome = re.sub(r"\b\d{1,2}[:h]\d{2}(?:\s*[-–]\s*\d{1,2}[:h]\d{2})?\b", "", nome, flags=re.IGNORECASE)
    if categoria:
        nome = re.sub(
            r"(?<!\w)(?:TIO|TID|TaIO|TaID|TCO|TCD|FIS|PSI)(?!\w)",
            "",
            nome,
            flags=re.IGNORECASE,
        )
        nome = re.sub(
            r"\b(?:t[eé]cnica|t[aá]tica)\s+(?:individual|coletiva)\s+(?:ofensiva|defensiva)\b"
            r"|\bprepara[cç][aã]o\s+f[ií]sica\b|\bpsicol[oó]gic[oa]\b",
            "",
            nome,
            flags=re.IGNORECASE,
        )
    if minutos is not None:
        nome = re.sub(rf"(?<!\d){minutos}\s*(?:min(?:uto)?s?|mn|')?", "", nome, flags=re.IGNORECASE)
    nome = re.sub(r"\b(?:bloco)\s*[A-D]\b", "", nome, flags=re.IGNORECASE)
    nome = re.sub(r"\s*[|;]+\s*", " - ", nome)
    return re.sub(r"^[\s:,.\-–]+|[\s:,.\-–]+$", "", re.sub(r"\s+", " ", nome))[:180]


def interpretar_texto(texto: str) -> dict:
    texto = (texto or "").replace("\r", "\n")
    linhas = [re.sub(r"\s+", " ", linha).strip() for linha in texto.split("\n") if linha.strip()]
    cabecalho = {
        "data": "",
        "hora_inicio": "",
        "hora_fim": "",
        "local": "",
        "escalao": "",
        "objectivos": "",
        "rpe_geral": "",
        "volume_minutos": "",
    }
    avisos = []

    data = re.search(r"\b(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{2,4})\b", texto)
    if data:
        ano = int(data.group(3))
        ano += 2000 if ano < 100 else 0
        try:
            data_valida = date(ano, int(data.group(2)), int(data.group(1)))
            cabecalho["data"] = data_valida.isoformat()
        except ValueError:
            avisos.append("A data do documento não é válida e deve ser confirmada manualmente.")
    intervalo = re.search(r"\b(\d{1,2})[:h](\d{2})\s*[-–aà]\s*(\d{1,2})[:h](\d{2})\b", texto, re.IGNORECASE)
    if intervalo:
        cabecalho["hora_inicio"] = f"{int(intervalo.group(1)):02d}:{intervalo.group(2)}"
        cabecalho["hora_fim"] = f"{int(intervalo.group(3)):02d}:{intervalo.group(4)}"

    for linha in linhas:
        chave = _sem_acentos(linha)
        if chave.startswith(("LOCAL ", "LOCAL:", "PAVILHAO ")):
            cabecalho["local"] = re.split(r"[:|]", linha, maxsplit=1)[-1].strip(" -")[:180]
        elif chave.startswith(("ESCALAO ", "ESCALAO:")):
            cabecalho["escalao"] = re.split(r"[:|]", linha, maxsplit=1)[-1].strip(" -")[:80]
        elif chave.startswith((
            "OBJETIVO ", "OBJETIVO:", "OBJETIVOS ", "OBJETIVOS:",
            "OBJECTIVO ", "OBJECTIVO:", "OBJECTIVOS ", "OBJECTIVOS:",
        )):
            cabecalho["objectivos"] = re.split(r"[:|]", linha, maxsplit=1)[-1].strip(" -")
        elif "RPE GERAL" in chave or "RPE SESSAO" in chave:
            cabecalho["rpe_geral"] = _primeiro_grupo(r"RPE(?:\s+(?:geral|sess[aã]o))?\s*[:\-]?\s*(10|[0-9](?:[,.]5)?)", linha).replace(",", ".")

    exercicios = []
    bloco_atual = "A"
    colunas_tabela: dict[str, int] = {}
    ignorar = {"DATA", "HORA", "MINUTOS", "EXERCICIO", "ORGANIZACAO", "COACHING POINTS", "CATEGORIA"}
    for linha in linhas:
        bloco = re.search(r"\bBLOCO\s*([A-D])\b", _sem_acentos(linha))
        if bloco and len(linha) <= 35:
            bloco_atual = bloco.group(1)
            continue
        if "|" in linha:
            celulas = [celula.strip() for celula in linha.split("|")]
            chaves_celulas = [_sem_acentos(celula).strip(" :") for celula in celulas]
            mapa_candidato = {}
            for indice, chave_celula in enumerate(chaves_celulas):
                if chave_celula in {"HORA", "INICIO"}:
                    mapa_candidato["hora"] = indice
                elif chave_celula in {"MINUTOS", "MIN", "DURACAO", "TEMPO"}:
                    mapa_candidato["minutos"] = indice
                elif chave_celula in {"EXERCICIO", "EXERCICIOS", "CONTEUDO", "TAREFA"}:
                    mapa_candidato["nome"] = indice
                elif chave_celula in {"CATEGORIA", "TIPO", "COMPONENTE"}:
                    mapa_candidato["categoria"] = indice
                elif chave_celula in {"RPE", "INTENSIDADE"}:
                    mapa_candidato["rpe"] = indice
                elif chave_celula.startswith("ORGANIZACAO"):
                    mapa_candidato["organizacao"] = indice
                elif chave_celula.startswith("COACHING") or chave_celula in {"PONTOS CHAVE", "CORRECOES"}:
                    mapa_candidato["coaching_points"] = indice
            if "nome" in mapa_candidato and "minutos" in mapa_candidato:
                colunas_tabela = mapa_candidato
                continue
            if colunas_tabela:
                obter = lambda chave: celulas[colunas_tabela[chave]].strip() if chave in colunas_tabela and colunas_tabela[chave] < len(celulas) else ""
                nome_celula = obter("nome")
                categoria = normalizar_categoria(obter("categoria") or nome_celula or linha)
                minutos_celula = _primeiro_grupo(r"(\d{1,3})", obter("minutos"))
                minutos = int(minutos_celula) if minutos_celula and 1 <= int(minutos_celula) <= 240 else None
                if minutos and nome_celula:
                    nome = _limpar_nome_exercicio(nome_celula, categoria, minutos)
                    rpe = _primeiro_grupo(r"(10|[0-9](?:[,.]5)?)", obter("rpe")).replace(",", ".")
                    exercicios.append({
                        "bloco": bloco_atual,
                        "nome": nome or f"Exercício {len(exercicios) + 1}",
                        "categoria": categoria or "",
                        "hora": _primeiro_grupo(r"(\d{1,2}[:h]\d{2})", obter("hora")).replace("h", ":"),
                        "minutos": minutos,
                        "rpe": rpe,
                        "organizacao": obter("organizacao")[:2000],
                        "coaching_points": obter("coaching_points")[:2000],
                        "confianca": "alta" if categoria else "media",
                    })
                    continue
        categoria = normalizar_categoria(linha)
        minutos = _extrair_minutos(linha, categoria)
        chave = _sem_acentos(linha).strip(" :|-")
        parece_exercicio = bool(
            minutos
            and chave not in ignorar
            and not chave.startswith(("VOLUME", "DURACAO TOTAL", "TEMPO TOTAL", "RPE GERAL"))
        )
        if not parece_exercicio:
            continue
        hora = _primeiro_grupo(r"\b(\d{1,2}[:h]\d{2})\b", linha).replace("h", ":")
        rpe = _primeiro_grupo(r"\bRPE\s*[:\-]?\s*(10|[0-9](?:[,.]5)?)\b", linha).replace(",", ".")
        nome = _limpar_nome_exercicio(linha, categoria, minutos)
        if not nome:
            nome = f"Exercício {len(exercicios) + 1}"
        exercicios.append({
            "bloco": bloco_atual,
            "nome": nome,
            "categoria": categoria or "",
            "hora": hora,
            "minutos": minutos or "",
            "rpe": rpe,
            "organizacao": "",
            "coaching_points": "",
            "confianca": "alta" if categoria and minutos and nome else "media",
        })

    total = sum(int(item["minutos"]) for item in exercicios if item["minutos"])
    cabecalho["volume_minutos"] = str(total) if total else ""
    if not exercicios:
        avisos.append("Não foram reconhecidos exercícios completos. Adiciona ou corrige as linhas antes de aplicar.")
    elif any(not item["categoria"] for item in exercicios):
        avisos.append("Alguns exercícios não têm categoria reconhecida; confirma-os antes de aplicar.")

    totais = {categoria: 0 for categoria in CATEGORIAS_EXERCICIO}
    for exercicio in exercicios:
        if exercicio["categoria"]:
            totais[exercicio["categoria"]] += int(exercicio["minutos"] or 0)
    return {"cabecalho": cabecalho, "exercicios": exercicios, "totais_categoria": totais, "avisos": avisos}


def sincronizar_minutos_categorias(plano, db=None) -> dict[str, int]:
    """Mantém o balanço da sessão coerente com os exercícios guardados."""
    totais = {categoria: 0 for categoria in CATEGORIAS_EXERCICIO}
    if db is not None:
        from models.treino import ExercicioPlano

        exercicios = db.query(ExercicioPlano).filter(ExercicioPlano.plano_id == plano.id).all()
    else:
        exercicios = plano.exercicios or []
    for exercicio in exercicios:
        categoria = normalizar_categoria(exercicio.categoria)
        if categoria:
            totais[categoria] += exercicio.minutos or 0
    sessao = plano.sessao
    for categoria, campo in CAMPOS_CATEGORIA.items():
        setattr(sessao, campo, totais[categoria])
    return totais
